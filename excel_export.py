"""Phase 5: professional .xlsx M&A model export (presentation layer only).

Consumes the Deal, the approved DealTeamResult (deterministic contexts + current
reviewed agent outputs), and optionally a MemoResult. It never calls an LLM, never
alters its inputs, and creates no new economics: every Excel formula mirrors an
existing valuation.py formula, and every formula is evaluated and checked against
the deterministic Python engine before the file is saved.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from agents.ic_memo_writer import PERMITTED_CONCLUSIONS
from agents.registry import COMMERCIAL_DD, FINANCIAL_DD, LEGAL_OPS_DD, RISK
from deal_analysis import (
    build_synergy_context,
    build_transaction_structure_context,
    build_valuation_context,
    lineage_from_contexts,
    strip_input_lineage,
)
from excel_formulas import FormulaError, WorkbookEvaluator
from ic_review import ICReview
from memo import (
    MemoGenerationError,
    MemoResult,
    check_memo_eligibility,
    current_agent_versions,
)
from models import Deal
from orchestrator import APPROVED, DealTeamResult

SHEET_ORDER: tuple[str, ...] = (
    "IC Summary", "Assumptions", "Valuation", "Synergies", "Transaction Structure",
    "Accretion Dilution", "Due Diligence", "Risk Summary", "IC Memo",
)
DEFAULT_OUTPUT = Path("reports/Northstar_Atlas_MA_Model.xlsx")
FOOTER_NOTE = "Generated from deterministic Python calculations and final reviewed AI analyses."
NO_MEMO_NOTE = ("Investment Committee memo has not been generated for this reviewed "
                "analytical state.")
FINANCING_CONVENTION = ("The model assumes the cash portion of the equity purchase price is "
                        "funded entirely with new debt at the supplied new-debt interest rate.")
FINANCING_WARNING = "This is a modelling convention, NOT a confirmed financing plan."
SYNERGY_CASE_CAUTION = ("Synergy-inclusive accretion assumes full supplied run-rate synergies "
                        "and should not be interpreted as the expected outcome.")
PYTHON_SOURCE = "Source: deterministic Python — valuation.py / deal_analysis.py"

# --- Styles ------------------------------------------------------------------------------

FONT_NAME = "Arial"
NAVY = "1F3864"
BLUE, BLACK, GREEN, GREY = "0000FF", "000000", "008000", "595959"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")

FMT_MONEY = "#,##0.0;[Red](#,##0.0);-"
FMT_MONEY_PRECISE = "#,##0.0000;[Red](#,##0.0000);-"
FMT_PCT = "0.0%;[Red](0.0%);-"
FMT_MULTIPLE = "0.0x"
FMT_EPS = "0.0000;[Red](0.0000);-"
FMT_SHARES = "#,##0.00"
FMT_OWNERSHIP = "0.0%"
FMT_FACTOR = "0.0000"
FMT_INT = "0"
FMT_PRICE = "#,##0.00"


def _font(color: str = BLACK, *, bold: bool = False, italic: bool = False,
          size: int = 10) -> Font:
    return Font(name=FONT_NAME, size=size, color=color, bold=bold, italic=italic)


class ExcelExportError(Exception):
    """The analytical state cannot be exported, or the workbook failed validation."""


# --- Eligibility ---------------------------------------------------------------------------

def check_export_eligibility(
    deal: Deal, result: DealTeamResult, memo_result: MemoResult | None
) -> ICReview:
    """Return the final approved IC review, or raise ExcelExportError."""
    try:
        final_review = check_memo_eligibility(result)  # same Phase-3 gate as the memo
    except MemoGenerationError as exc:
        raise ExcelExportError(f"cannot export this analytical state: {exc}") from exc

    expected = {
        "valuation": build_valuation_context(deal),
        "synergy": build_synergy_context(deal),
        "transaction_structure": build_transaction_structure_context(deal),
    }
    for key, context in expected.items():
        actual = result.deterministic_contexts.get(key)
        # input-lineage notes are metadata; every calculated value must match exactly
        if actual is None or strip_input_lineage(actual) != strip_input_lineage(context):
            raise ExcelExportError(
                f"the result's {key} context does not match the deal's deterministic output")

    if memo_result is not None:
        if memo_result.orchestration_status != APPROVED:
            raise ExcelExportError("the memo was not generated from an approved review")
        if memo_result.deal_name != deal.deal_name:
            raise ExcelExportError("the memo belongs to a different deal")
        if memo_result.conclusion not in PERMITTED_CONCLUSIONS:
            raise ExcelExportError("the memo conclusion is not a permitted conclusion")
        if (memo_result.generated_from_review_round != result.review_history[-1].round
                or memo_result.review_rounds_used != result.review_rounds_used
                or memo_result.source_agent_versions != current_agent_versions(result)):
            raise ExcelExportError(
                "the memo does not correspond to this reviewed result (different review "
                "round or agent versions)")
    return final_review


# --- Sheet builder -------------------------------------------------------------------------

@dataclass
class _Model:
    """Workbook-wide registry of named cells and expected values to check."""

    workbook: Workbook
    refs: dict[str, tuple[str, str]] = field(default_factory=dict)
    expectations: list[tuple[str, str, Any]] = field(default_factory=list)


class _Sheet:
    def __init__(self, model: _Model, title: str, widths: dict[str, float]):
        self.model = model
        self.ws: Worksheet = model.workbook[title]
        self.row = 1
        self.ws.sheet_view.showGridLines = False
        for column, width in widths.items():
            self.ws.column_dimensions[column].width = width
        self.text_width = widths.get("A", 60)

    # references ------------------------------------------------------------------------
    def ref(self, key: str) -> str:
        sheet, coordinate = self.model.refs[key]
        column = "".join(ch for ch in coordinate if ch.isalpha())
        row = coordinate[len(column):]
        absolute = f"${column}${row}"
        return absolute if sheet == self.ws.title else f"'{sheet}'!{absolute}"

    def register(self, key: str, coordinate: str) -> None:
        if key in self.model.refs:
            raise AssertionError(f"duplicate cell key {key}")
        self.model.refs[key] = (self.ws.title, coordinate)

    # layout --------------------------------------------------------------------------------
    def title(self, text: str) -> None:
        self.ws.cell(1, 1, text).font = _font(NAVY, bold=True, size=14)
        self.ws.cell(2, 1, FOOTER_NOTE).font = _font(GREY, italic=True, size=9)
        self.ws.oddFooter.center.text = FOOTER_NOTE
        self.row = 4

    def section(self, text: str, columns: int = 3) -> None:
        for col in range(1, columns + 1):
            cell = self.ws.cell(self.row, col, text if col == 1 else None)
            cell.fill = HEADER_FILL
            cell.font = _font("FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="left", vertical="center")
        self.row += 1

    def header(self, *labels: str) -> None:
        for col, label in enumerate(labels, start=1):
            cell = self.ws.cell(self.row, col, label)
            cell.font = _font(bold=True)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(horizontal="left" if col == 1 else "right",
                                       wrap_text=True)
        self.row += 1

    def skip(self, rows: int = 1) -> None:
        self.row += rows

    def put(self, column: int, value: Any, *, fmt: str | None = None, key: str | None = None,
            input_source: str | None = None, comment: str | None = None,
            unverified: bool = False, expect: Any = None, bold: bool = False) -> str:
        """Write one value cell. Font colour follows the finance convention: inputs
        blue, same-sheet formulas black, cross-sheet formulas green."""
        cell = self.ws.cell(self.row, column, value)
        is_formula = isinstance(value, str) and value.startswith("=")
        if isinstance(value, str) and not is_formula:
            cell.data_type = "s"
        if input_source is not None:
            color = BLUE
            comment = comment or f"Source: {input_source}"
        elif is_formula and "!" in value:
            color = GREEN
        else:
            color = BLACK
        cell.font = _font(color, bold=bold)
        cell.alignment = Alignment(horizontal="right" if column > 1 else "left",
                                   vertical="top")
        if fmt:
            cell.number_format = fmt
        if unverified:
            cell.fill = YELLOW_FILL
        if comment:
            cell.comment = Comment(comment, "Deal Team Model", width=320, height=90)
        if key:
            self.register(key, cell.coordinate)
        if expect is not None:
            self.model.expectations.append((self.ws.title, cell.coordinate, expect))
        return cell.coordinate

    def item(self, label: str, value: Any, *, note: str | None = None, **kwargs) -> str:
        label_cell = self.ws.cell(self.row, 1, label)
        label_cell.font = _font()
        label_cell.alignment = Alignment(horizontal="left", vertical="top")
        coordinate = self.put(2, value, **kwargs)
        if note:
            note_cell = self.ws.cell(self.row, 3, note)
            note_cell.data_type = "s"
            note_cell.font = _font(GREY, italic=True)
            note_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        self.row += 1
        return coordinate

    def text(self, text: str, *, bold: bool = False, italic: bool = False,
             color: str = BLACK, size: int = 10, comment: str | None = None,
             fill: PatternFill | None = None, border: bool = False) -> None:
        cell = self.ws.cell(self.row, 1, text)
        cell.data_type = "s"  # never let agent text become a formula
        cell.font = _font(color, bold=bold, italic=italic, size=size)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if fill:
            cell.fill = fill
        if border:
            cell.border = Border(bottom=Side(style="thin", color=NAVY))
        if comment:
            cell.comment = Comment(comment, "Deal Team Model", width=320, height=90)
        chars_per_line = max(20, int(self.text_width * 1.2))
        lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in text.split("\n"))
        if lines > 1:
            self.ws.row_dimensions[self.row].height = 13.5 * lines
        self.row += 1

    def markdown(self, text: str, source_comment: str) -> None:
        """Render Markdown lines with visual hierarchy only; content is not rewritten."""
        first = True
        for line in text.splitlines():
            if not line.strip():
                continue
            comment = source_comment if first else None
            first = False
            stripped = line.strip()
            if stripped.startswith("# "):
                self.text(stripped[2:], bold=True, color=NAVY, size=12, comment=comment)
            elif stripped.startswith("## "):
                self.skip()
                self.text(stripped[3:], bold=True, color=NAVY, comment=comment, border=True)
            elif stripped.startswith("### "):
                self.text(stripped[4:], bold=True, comment=comment)
            else:
                self.text(line.rstrip(), comment=comment)

    def freeze(self, cell: str = "A4") -> None:
        self.ws.freeze_panes = cell

    def print_setup(self) -> None:
        ps = self.ws.page_setup
        ps.orientation = "landscape"
        ps.fitToWidth = 1
        ps.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.page_margins.left = self.ws.page_margins.right = 0.5
        self.ws.page_margins.top = self.ws.page_margins.bottom = 0.6
        self.ws.print_title_rows = "1:2"


# --- Sheet contents --------------------------------------------------------------------------

def _money_label(label: str, deal: Deal) -> str:
    return f"{label} ({deal.currency} {deal.units})"


def _build_assumptions(model: _Model, deal: Deal, source: str,
                       input_sources: Mapping[str, str] | None = None,
                       share_basis: str | None = None,
                       data_quality: Sequence[str] = ()) -> None:
    """Input cells hold the EFFECTIVE deal values. Cell comments carry lineage: source,
    any override (value, as-of, source, reason) and the original researched value."""
    s = _Sheet(model, "Assumptions", {"A": 46, "B": 18, "C": 62})
    s.title("Assumptions — Deal Inputs")
    s.text("Blue font = supplied input. Yellow fill = unverified assumption requiring "
           "validation. All monetary values in " f"{deal.currency} {deal.units}.",
           italic=True, color=GREY)
    s.skip()

    def inp(label, key, value, fmt=None, unverified=False, note=None):
        cell_source = (input_sources[key] if input_sources and key in input_sources
                       else f"{source} — {key}")
        if key == "acquirer.shares_outstanding" and share_basis:
            cell_source += f" Share basis definition: {share_basis}."
        if data_quality and key.endswith((".cash", ".debt", ".shares_outstanding")):
            cell_source += " Data-quality note: " + " ".join(data_quality)
        s.item(label, value, fmt=fmt, key=key, input_source=cell_source,
               unverified=unverified, note=note)

    a, t, m = deal.acquirer, deal.target, deal.market_data
    d, y, x = deal.dcf_assumptions, deal.synergy_assumptions, deal.transaction_assumptions
    s.section("DEAL OVERVIEW")
    inp("Deal name", "deal_name", deal.deal_name)
    inp("Currency", "currency", deal.currency)
    inp("Units", "units", deal.units)
    s.skip()
    s.section("ACQUIRER FINANCIALS")
    inp("Acquirer", "acquirer.name", a.name)
    inp(_money_label("Revenue", deal), "acquirer.revenue", a.revenue, FMT_MONEY)
    inp(_money_label("EBITDA", deal), "acquirer.ebitda", a.ebitda, FMT_MONEY)
    inp(_money_label("Net income", deal), "acquirer.net_income", a.net_income, FMT_MONEY)
    inp(_money_label("Cash", deal), "acquirer.cash", a.cash, FMT_MONEY)
    inp(_money_label("Debt", deal), "acquirer.debt", a.debt, FMT_MONEY)
    inp(f"Acquirer share basis ({deal.units} shares)", "acquirer.shares_outstanding",
        a.shares_outstanding, FMT_SHARES, note=share_basis)
    inp(f"Share price ({deal.currency} per share)", "acquirer.share_price", a.share_price,
        FMT_PRICE)
    s.skip()
    s.section("TARGET FINANCIALS")
    inp("Target", "target.name", t.name)
    inp(_money_label("Revenue", deal), "target.revenue", t.revenue, FMT_MONEY)
    inp(_money_label("EBITDA", deal), "target.ebitda", t.ebitda, FMT_MONEY)
    inp(_money_label("Net income", deal), "target.net_income", t.net_income, FMT_MONEY)
    inp(_money_label("Free cash flow", deal), "target.free_cash_flow", t.free_cash_flow,
        FMT_MONEY)
    inp(_money_label("Cash", deal), "target.cash", t.cash, FMT_MONEY)
    inp(_money_label("Debt", deal), "target.debt", t.debt, FMT_MONEY)
    s.skip()
    s.section("MARKET VALUATION INPUTS")
    for label, key, value, optional in (
        ("Comparable EV / EBITDA (x)", "market_data.comparable_ev_ebitda_multiple",
         m.comparable_ev_ebitda_multiple, False),
        ("Precedent EV / EBITDA (x)", "market_data.precedent_ev_ebitda_multiple",
         m.precedent_ev_ebitda_multiple, False),
        ("Comparable EV / Revenue (x)", "market_data.comparable_ev_revenue_multiple",
         m.comparable_ev_revenue_multiple, True),
        ("Precedent EV / Revenue (x)", "market_data.precedent_ev_revenue_multiple",
         m.precedent_ev_revenue_multiple, True),
    ):
        if value is None and optional:
            continue  # optional multiple not supplied: no row
        inp(label, key, value if value is not None else "Not supplied",
            FMT_MULTIPLE if value is not None else None)
    s.skip()
    s.section("DCF ASSUMPTIONS")
    inp("Forecast years", "dcf_assumptions.forecast_years", d.forecast_years, FMT_INT)
    inp("FCF growth rate", "dcf_assumptions.fcf_growth_rate", d.fcf_growth_rate, FMT_PCT,
        unverified=True)
    inp("WACC", "dcf_assumptions.wacc", d.wacc, FMT_PCT, unverified=True)
    inp("Terminal growth rate", "dcf_assumptions.terminal_growth_rate",
        d.terminal_growth_rate, FMT_PCT, unverified=True)
    s.skip()
    s.section("SYNERGY ASSUMPTIONS (UNVALIDATED)")
    inp(_money_label("Annual cost synergy", deal), "synergy_assumptions.annual_cost_synergy",
        y.annual_cost_synergy, FMT_MONEY, unverified=True)
    inp(_money_label("Annual revenue synergy EBITDA", deal),
        "synergy_assumptions.annual_revenue_synergy_ebitda",
        y.annual_revenue_synergy_ebitda, FMT_MONEY, unverified=True)
    inp("Realization years", "synergy_assumptions.realization_years", y.realization_years,
        FMT_INT, unverified=True)
    inp("Synergy tax rate", "synergy_assumptions.tax_rate", y.tax_rate, FMT_PCT,
        unverified=True)
    inp("Synergy discount rate", "synergy_assumptions.discount_rate", y.discount_rate,
        FMT_PCT, unverified=True)
    inp(_money_label("Integration cost", deal), "synergy_assumptions.integration_cost",
        y.integration_cost, FMT_MONEY, unverified=True)
    s.skip()
    s.section("TRANSACTION ASSUMPTIONS")
    inp(_money_label("Offer enterprise value", deal),
        "transaction_assumptions.offer_enterprise_value", x.offer_enterprise_value, FMT_MONEY)
    inp("Cash percentage", "transaction_assumptions.cash_percentage", x.cash_percentage,
        FMT_PCT)
    inp("Stock percentage", "transaction_assumptions.stock_percentage", x.stock_percentage,
        FMT_PCT)
    inp("New debt interest rate", "transaction_assumptions.new_debt_interest_rate",
        x.new_debt_interest_rate, FMT_PCT, unverified=True,
        note="Financing modelling convention input; not a confirmed financing term.")
    inp("Transaction tax rate", "transaction_assumptions.tax_rate", x.tax_rate, FMT_PCT)
    s.freeze("A4")
    s.print_setup()


# Valuation methods: context key, summary method name, sheet key, section title, metric.
METHOD_SPECS = (
    ("comparable_company", "Comparable Companies", "val.comps_ev", "COMPARABLE COMPANIES",
     "ebitda"),
    ("precedent_transaction", "Precedent Transactions", "val.precedent_ev",
     "PRECEDENT TRANSACTIONS", "ebitda"),
    ("comparable_company_revenue", "Comparable Companies (EV/Revenue)", "val.comps_rev_ev",
     "COMPARABLE COMPANIES — EV / REVENUE", "revenue"),
    ("precedent_transaction_revenue", "Precedent Transactions (EV/Revenue)",
     "val.precedent_rev_ev", "PRECEDENT TRANSACTIONS — EV / REVENUE", "revenue"),
)
METHOD_KEYS = {name: key for _, name, key, _, _ in METHOD_SPECS}
METHOD_KEYS["Discounted Cash Flow"] = "val.dcf_ev"
MULTIPLE_PATHS = {
    ("comparable_company", "ebitda"): ("Comparable EV / EBITDA (x)",
                                        "market_data.comparable_ev_ebitda_multiple"),
    ("precedent_transaction", "ebitda"): ("Precedent EV / EBITDA (x)",
                                           "market_data.precedent_ev_ebitda_multiple"),
    ("comparable_company_revenue", "revenue"): ("Comparable EV / Revenue (x)",
                                                "market_data.comparable_ev_revenue_multiple"),
    ("precedent_transaction_revenue", "revenue"): ("Precedent EV / Revenue (x)",
                                                   "market_data.precedent_ev_revenue_multiple"),
}
NOT_APPLICABLE = "NOT APPLICABLE"


def _short_method(name: str) -> str:
    return "DCF" if name == "Discounted Cash Flow" else name


def _cell_list(rows: list[int]) -> str:
    if rows == list(range(rows[0], rows[-1] + 1)):
        return f"B{rows[0]}:B{rows[-1]}"
    return ",".join(f"B{r}" for r in rows)


def _build_valuation(model: _Model, deal: Deal, val: dict, labels: dict) -> None:
    s = _Sheet(model, "Valuation", {"A": 38, "B": 18, "C": 18, "D": 18, "E": 34, "F": 16})
    s.title("Valuation Analysis")
    money = _money_label("Enterprise value", deal)
    summary = val["valuation_summary"]
    methods = summary["methods"]

    # Detail sections first (below the summary) so the summary can reference them.
    s.row = max(19, 16 + len(methods))
    for context_key, name, key, title, metric in METHOD_SPECS:
        if context_key not in val:
            continue  # optional EV/Revenue method not supplied
        method = val[context_key]
        s.section(title, 4)
        label, path = MULTIPLE_PATHS[(context_key, metric)]
        metric_label = "Target EBITDA" if metric == "ebitda" else "Target revenue"
        metric_cell = s.item(_money_label(metric_label, deal), f"={s.ref('target.' + metric)}",
                             fmt=FMT_MONEY)
        if not method["applicable"]:
            s.item(f"Implied {money}", NOT_APPLICABLE, bold=True, note=method["reason"],
                   comment=PYTHON_SOURCE)
            s.skip()
            continue
        multiple = s.item(label, f"={s.ref(path)}", fmt=FMT_MULTIPLE)
        s.item(f"Implied {money}", f"={metric_cell}*{multiple}", fmt=FMT_MONEY, key=key,
               expect=method["enterprise_value"])
        s.skip()

    dcf = val["dcf"]
    s.section("DISCOUNTED CASH FLOW", 4)
    base = s.item(_money_label("Base free cash flow", deal),
                  f"={s.ref('target.free_cash_flow')}", fmt=FMT_MONEY)
    growth = s.item("FCF growth rate", f"={s.ref('dcf_assumptions.fcf_growth_rate')}",
                    fmt=FMT_PCT)
    wacc = s.item("WACC", f"={s.ref('dcf_assumptions.wacc')}", fmt=FMT_PCT)
    tg = s.item("Terminal growth rate", f"={s.ref('dcf_assumptions.terminal_growth_rate')}",
                fmt=FMT_PCT)
    s.item("Forecast years", f"={s.ref('dcf_assumptions.forecast_years')}", fmt=FMT_INT)
    s.skip()
    s.header("Year", f"FCF ({deal.currency} {deal.units})", "Discount factor",
             f"Discounted FCF ({deal.currency} {deal.units})")
    first = s.row
    for i, year in enumerate(dcf["years"]):
        s.put(1, year, fmt=FMT_INT,
              comment="Schedule period from deterministic Python (forecast_years)"
              if i == 0 else None)
        s.put(2, f"={base}*(1+{growth})^A{s.row}", fmt=FMT_MONEY_PRECISE,
              expect=dcf["forecast_fcfs"][i])
        s.put(3, f"=1/(1+{wacc})^A{s.row}", fmt=FMT_FACTOR, expect=dcf["discount_factors"][i])
        s.put(4, f"=B{s.row}*C{s.row}", fmt=FMT_MONEY_PRECISE, expect=dcf["discounted_fcfs"][i])
        s.row += 1
    last = s.row - 1
    s.skip()
    pv = s.item(_money_label("PV of forecast FCF", deal), f"=SUM(D{first}:D{last})",
                fmt=FMT_MONEY_PRECISE, expect=dcf["pv_forecast_fcfs"])
    tv = s.item(_money_label("Terminal value", deal),
                f"=B{last}*(1+{tg})/({wacc}-{tg})", fmt=FMT_MONEY_PRECISE,
                expect=dcf["terminal_value"])
    dtv = s.item(_money_label("Discounted terminal value", deal), f"={tv}*C{last}",
                 fmt=FMT_MONEY_PRECISE, expect=dcf["discounted_terminal_value"])
    ev = s.item(f"DCF {money}", f"={pv}+{dtv}", fmt=FMT_MONEY_PRECISE, key="val.dcf_ev",
                expect=dcf["enterprise_value"], bold=True)
    s.item("Terminal value share of EV", f"={dtv}/{ev}", fmt=FMT_PCT,
           expect=dcf["terminal_value_share_of_ev"])
    bottom = s.row + 1

    # Summary at the top: every method, with applicability, from the Python summary.
    s.row = 4
    s.section("VALUATION SUMMARY", 3)
    s.header("Method", money, "Offer EV vs method")
    in_range = []
    for method in methods:
        name = method["method"]
        if not method["applicable"]:
            s.put(3, method["reason"], comment=PYTHON_SOURCE)
            s.item(name, NOT_APPLICABLE, comment=PYTHON_SOURCE)
            continue
        if method["included_in_range"]:
            in_range.append(s.row)
            s.put(3, labels["offer_vs_method"][name], comment=PYTHON_SOURCE)
            s.item(name, f"={s.ref(METHOD_KEYS[name])}", fmt=FMT_MONEY,
                   key=f"val.summary.{name}", expect=summary["by_method"][name])
        else:
            s.put(3, "Supplementary — not in the valuation range", comment=PYTHON_SOURCE)
            s.item(name, f"={s.ref(METHOD_KEYS[name])}", fmt=FMT_MONEY,
                   key=f"val.summary.{name}", expect=method["enterprise_value"])
    if summary["range_available"]:
        cells = _cell_list(in_range)
        low = s.item("Low", f"=MIN({cells})", fmt=FMT_MONEY, key="val.low",
                     expect=summary["low"])
        high = s.item("High", f"=MAX({cells})", fmt=FMT_MONEY, key="val.high",
                      expect=summary["high"])
        s.item("Midpoint", f"=({low}+{high})/2", fmt=FMT_MONEY, expect=summary["midpoint"])
        s.item("Average", f"=AVERAGE({cells})", fmt=FMT_MONEY, expect=summary["average"])
    else:
        for label in ("Low", "High", "Midpoint", "Average"):
            s.item(label, "Not available", note=summary["range_note"] if label == "Low" else None,
                   comment=PYTHON_SOURCE)
    s.put(3, f"Offer vs range: {labels['offer_vs_range']}", comment=PYTHON_SOURCE)
    s.item("Offer Enterprise Value",
           f"={s.ref('transaction_assumptions.offer_enterprise_value')}",
           fmt=FMT_MONEY, key="val.offer_ev", bold=True)
    if summary["range_available"]:
        s.text(f"Low method: {summary['low_method']}; high method: {summary['high_method']}. "
               "Relationship labels are deterministic Python outputs; no premium or discount "
               "is calculated.", italic=True, color=GREY)
    else:
        s.text(summary["range_note"], italic=True, color=GREY)

    # Chart data (range methods + offer) and chart.
    s.row = bottom
    s.section("CHART DATA — ENTERPRISE VALUE COMPARISON", 2)
    chart_first = s.row
    chart_items = [(_short_method(m["method"]), METHOD_KEYS[m["method"]]) for m in methods
                   if m["applicable"] and m["included_in_range"]]
    chart_items.append(("Offer EV", "val.offer_ev"))
    for label, key in chart_items:
        s.item(label, f"={s.ref(key)}", fmt=FMT_MONEY)
    chart_last = chart_first + len(chart_items) - 1
    chart = BarChart()
    chart.type = "col"
    chart.title = f"Enterprise value comparison ({deal.currency} {deal.units})"
    chart.y_axis.title = f"{deal.currency} {deal.units}"
    chart.legend = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.add_data(Reference(s.ws, min_col=2, min_row=chart_first, max_row=chart_last))
    chart.set_categories(Reference(s.ws, min_col=1, min_row=chart_first, max_row=chart_last))
    chart.height, chart.width = 7.5, 15
    s.ws.add_chart(chart, "E4")
    s.freeze("A4")
    s.print_setup()


def _build_synergies(model: _Model, deal: Deal, syn: dict, labels: dict) -> None:
    s = _Sheet(model, "Synergies", {"A": 42, "B": 18, "C": 18, "D": 22})
    s.title("Synergy Analysis")
    s.section("SYNERGY ASSUMPTIONS (UNVALIDATED — LINKED TO ASSUMPTIONS)", 4)
    cost = s.item(_money_label("Annual cost synergy", deal),
                  f"={s.ref('synergy_assumptions.annual_cost_synergy')}", fmt=FMT_MONEY,
                  unverified=True)
    rev = s.item(_money_label("Annual revenue synergy EBITDA", deal),
                 f"={s.ref('synergy_assumptions.annual_revenue_synergy_ebitda')}",
                 fmt=FMT_MONEY, unverified=True)
    s.item("Realization years", f"={s.ref('synergy_assumptions.realization_years')}",
           fmt=FMT_INT, unverified=True)
    tax = s.item("Tax rate", f"={s.ref('synergy_assumptions.tax_rate')}", fmt=FMT_PCT)
    rate = s.item("Discount rate", f"={s.ref('synergy_assumptions.discount_rate')}",
                  fmt=FMT_PCT)
    s.skip()

    calc = syn["calculation"]
    s.section("SYNERGY VALUATION", 4)
    s.item(_money_label("Annual cost synergy", deal), f"={cost}", fmt=FMT_MONEY)
    s.item(_money_label("Annual revenue synergy EBITDA", deal), f"={rev}", fmt=FMT_MONEY)
    pre = s.item(_money_label("Annual pre-tax synergies", deal), f"={cost}+{rev}",
                 fmt=FMT_MONEY_PRECISE, expect=calc["annual_pre_tax_synergies"])
    after = s.item(_money_label("Annual after-tax synergies", deal), f"={pre}*(1-{tax})",
                   fmt=FMT_MONEY_PRECISE, key="syn.after_tax",
                   expect=calc["annual_after_tax_synergies"])
    pv_row = s.row
    s.skip()  # PV row filled after the schedule
    integ = s.item(_money_label("Integration cost (at close)", deal),
                   f"={s.ref('synergy_assumptions.integration_cost')}", fmt=FMT_MONEY,
                   unverified=True, key="syn.integration_cost")
    net_row = s.row
    s.skip(2)

    s.section("DISCOUNTED SYNERGY SCHEDULE", 4)
    s.header("Year", f"After-tax synergy ({deal.currency} {deal.units})", "Discount factor",
             f"Discounted synergy ({deal.currency} {deal.units})")
    first = s.row
    for i, year in enumerate(range(1, calc["realization_years"] + 1)):
        s.put(1, year, fmt=FMT_INT,
              comment="Schedule period from deterministic Python (realization_years)"
              if i == 0 else None)
        s.put(2, f"={after}", fmt=FMT_MONEY_PRECISE)
        s.put(3, f"=1/(1+{rate})^A{s.row}", fmt=FMT_FACTOR)
        s.put(4, f"=B{s.row}*C{s.row}", fmt=FMT_MONEY_PRECISE,
              expect=calc["discounted_synergies"][i])
        s.row += 1
    last = s.row - 1
    end = s.row

    s.row = pv_row
    pv = s.item(_money_label("PV of synergies", deal), f"=SUM(D{first}:D{last})",
                fmt=FMT_MONEY_PRECISE, key="syn.pv", expect=calc["pv_synergies"])
    s.row = net_row
    s.put(3, f"Status: {labels['net_synergy_value_status']}", comment=PYTHON_SOURCE)
    s.item(_money_label("Net synergy value", deal), f"={pv}-{integ}", fmt=FMT_MONEY_PRECISE,
           key="syn.net", expect=calc["net_synergy_value"], bold=True)
    s.row = end + 1

    s.section("MODELLING CONVENTIONS", 4)
    for line in [
        "Full run-rate synergies from year one (no ramp-up).",
        "Finite realization period; no terminal synergy value.",
        "Integration cost is incurred at close, undiscounted and not tax-deductible.",
        "Synergy assumptions are unvalidated DEAL DATA assumptions.",
        f"Integration cost vs PV of synergies (deterministic label): "
        f"{labels['integration_cost_vs_pv_synergies']}.",
    ]:
        s.text(f"• {line}")
    s.freeze("A4")
    s.print_setup()


def _build_transaction(model: _Model, deal: Deal, tx: dict) -> None:
    s = _Sheet(model, "Transaction Structure", {"A": 46, "B": 18, "C": 60})
    s.title("Transaction Structure")
    calc = tx["calculation"]
    s.section("PURCHASE PRICE")
    ev = s.item(_money_label("Offer enterprise value", deal),
                f"={s.ref('transaction_assumptions.offer_enterprise_value')}", fmt=FMT_MONEY)
    debt = s.item(_money_label("Target debt", deal), f"={s.ref('target.debt')}", fmt=FMT_MONEY)
    cash = s.item(_money_label("Target cash", deal), f"={s.ref('target.cash')}", fmt=FMT_MONEY)
    net_debt = s.item(_money_label("Target net debt", deal), f"={debt}-{cash}", fmt=FMT_MONEY,
                      key="tx.net_debt", expect=calc["target_net_debt"])
    equity = s.item(_money_label("Offer equity value", deal), f"={ev}-{net_debt}",
                    fmt=FMT_MONEY, key="tx.equity", expect=calc["offer_equity_value"], bold=True)
    s.skip()
    s.section("CONSIDERATION")
    cash_pct = s.item("Cash %", f"={s.ref('transaction_assumptions.cash_percentage')}",
                      fmt=FMT_PCT, key="tx.cash_pct")
    stock_pct = s.item("Stock %", f"={s.ref('transaction_assumptions.stock_percentage')}",
                       fmt=FMT_PCT, key="tx.stock_pct")
    cash_cons = s.item(_money_label("Cash consideration", deal), f"={equity}*{cash_pct}",
                       fmt=FMT_MONEY, key="tx.cash_consideration",
                       expect=calc["cash_consideration"])
    stock_cons = s.item(_money_label("Stock consideration", deal), f"={equity}*{stock_pct}",
                        fmt=FMT_MONEY, key="tx.stock_consideration",
                        expect=calc["stock_consideration"])
    s.skip()
    s.section("FINANCING MODEL CONVENTION")
    s.text(FINANCING_CONVENTION)
    s.text(FINANCING_WARNING, bold=True, fill=YELLOW_FILL)
    new_debt = s.item(_money_label("New debt under model convention", deal), f"={cash_cons}",
                      fmt=FMT_MONEY, expect=calc["new_debt_under_financing_convention"])
    rate = s.item("New debt interest rate",
                  f"={s.ref('transaction_assumptions.new_debt_interest_rate')}", fmt=FMT_PCT,
                  unverified=True)
    pre = s.item(_money_label("Pre-tax interest expense", deal), f"={new_debt}*{rate}",
                 fmt=FMT_MONEY_PRECISE, expect=calc["annual_pre_tax_interest_expense"])
    tax = s.item("Tax rate", f"={s.ref('transaction_assumptions.tax_rate')}", fmt=FMT_PCT)
    s.item(_money_label("After-tax interest expense", deal), f"={pre}*(1-{tax})",
           fmt=FMT_MONEY_PRECISE, key="tx.after_tax_interest",
           expect=calc["annual_after_tax_interest_expense"])
    s.text("Not modelled: leverage, debt capacity, covenant headroom, or credit-rating impact.",
           italic=True, color=GREY)
    s.skip()
    s.section("SHARE ISSUANCE / OWNERSHIP")
    price = s.item(f"Acquirer share price ({deal.currency} per share)",
                   f"={s.ref('acquirer.share_price')}", fmt=FMT_PRICE)
    existing = s.item(f"Existing shares ({deal.units})",
                      f"={s.ref('acquirer.shares_outstanding')}", fmt=FMT_SHARES)
    new = s.item(f"New shares issued ({deal.units})", f"={stock_cons}/{price}", fmt=FMT_SHARES,
                 key="tx.new_shares", expect=calc["new_shares_issued"])
    pro_forma = s.item(f"Pro forma shares ({deal.units})", f"={existing}+{new}",
                       fmt=FMT_SHARES, key="tx.pro_forma_shares",
                       expect=calc["pro_forma_shares_outstanding"])
    seller = s.item("Seller ownership", f"={new}/{pro_forma}", fmt=FMT_OWNERSHIP,
                    key="tx.seller_ownership", expect=calc["seller_ownership_of_pro_forma_shares"])
    s.item("Existing acquirer holder ownership", f"=1-{seller}", fmt=FMT_OWNERSHIP,
           expect=calc["existing_acquirer_holders_ownership_of_pro_forma_shares"])
    s.freeze("A4")
    s.print_setup()


def _build_accretion(model: _Model, deal: Deal, tx: dict) -> None:
    s = _Sheet(model, "Accretion Dilution", {"A": 46, "B": 18, "C": 60})
    s.title("Accretion / Dilution")
    cases = (
        ("A", "CASE A — EXCLUDING SYNERGIES", "accretion_dilution_excluding_synergies", None),
        ("B", "CASE B — INCLUDING FULL RUN-RATE SYNERGIES",
         "accretion_dilution_including_run_rate_synergies", "syn.after_tax"),
    )
    for case, heading, calc_key, synergy_key in cases:
        calc = tx["calculation"][calc_key]
        s.section(heading)
        if synergy_key:
            s.text(SYNERGY_CASE_CAUTION, bold=True, fill=YELLOW_FILL)
        acq = s.item(_money_label("Acquirer standalone net income", deal),
                     f"={s.ref('acquirer.net_income')}", fmt=FMT_MONEY)
        tgt = s.item(_money_label("Target net income", deal), f"={s.ref('target.net_income')}",
                     fmt=FMT_MONEY)
        interest = s.item(_money_label("After-tax interest expense", deal),
                          f"={s.ref('tx.after_tax_interest')}", fmt=FMT_MONEY_PRECISE)
        if synergy_key:
            syn = s.item(_money_label("Synergy contribution (after tax)", deal),
                         f"={s.ref(synergy_key)}", fmt=FMT_MONEY_PRECISE, unverified=True)
        else:
            syn = s.item(_money_label("Synergy contribution", deal), 0,
                         fmt=FMT_MONEY_PRECISE, comment="Case definition: synergies excluded")
        pf_ni = s.item(_money_label("Pro forma net income", deal),
                       f"={acq}+{tgt}-{interest}+{syn}", fmt=FMT_MONEY_PRECISE,
                       expect=calc["pro_forma_net_income"])
        existing = s.item(f"Existing shares ({deal.units})",
                          f"={s.ref('acquirer.shares_outstanding')}", fmt=FMT_SHARES)
        new = s.item(f"New shares issued ({deal.units})", f"={s.ref('tx.new_shares')}",
                     fmt=FMT_SHARES)
        pf_shares = s.item(f"Pro forma shares ({deal.units})", f"={existing}+{new}",
                           fmt=FMT_SHARES, expect=calc["pro_forma_shares"])
        std = s.item(f"Standalone EPS ({deal.currency} per share)", f"={acq}/{existing}",
                     fmt=FMT_EPS, expect=calc["standalone_eps"])
        pf = s.item(f"Pro forma EPS ({deal.currency} per share)", f"={pf_ni}/{pf_shares}",
                    fmt=FMT_EPS, expect=calc["pro_forma_eps"])
        change = s.item(f"EPS change ({deal.currency} per share)", f"={pf}-{std}",
                        fmt=FMT_EPS, expect=calc["eps_change"])
        s.item("EPS change %", f'=IF({std}=0,"n/a",{change}/ABS({std}))', fmt=FMT_PCT,
               expect=calc["eps_change_pct"] if calc["eps_change_pct"] is not None else "n/a")
        s.item("Result", f'=IF(ABS({change})<1E-12,"Neutral",IF({change}>0,"Accretive",'
                         f'"Dilutive"))', key=f"acc.result.{case}", expect=calc["result"],
               bold=True)
        s.skip()
    s.text("Both cases are single-year, run-rate EPS views from the deterministic "
           "accretion_dilution model. Neither case is a base, expected, or most likely case.",
           italic=True, color=GREY)
    s.freeze("A4")
    s.print_setup()


def _version_line(result: DealTeamResult, agent: str) -> str:
    ref = next(r for r in current_agent_versions(result) if r.agent == agent)
    when = ("initial analysis" if ref.after_review_round == 0
            else f"after IC review round {ref.after_review_round}")
    return f"Current version: {ref.version} — run kind: {ref.kind} ({when})"


def _build_due_diligence(model: _Model, result: DealTeamResult) -> None:
    s = _Sheet(model, "Due Diligence", {"A": 120})
    s.title("Due Diligence — Final Reviewed Outputs")
    for heading, agent in (("FINANCIAL DUE DILIGENCE", FINANCIAL_DD),
                           ("COMMERCIAL DUE DILIGENCE", COMMERCIAL_DD),
                           ("LEGAL & OPERATIONAL DUE DILIGENCE", LEGAL_OPS_DD)):
        s.section(heading, 1)
        s.text(_version_line(result, agent), italic=True, color=GREY)
        s.markdown(result.outputs[agent], f"Source: final reviewed output — {agent}")
        s.skip()
    s.print_setup()


def _build_risk(model: _Model, result: DealTeamResult, review: ICReview) -> None:
    s = _Sheet(model, "Risk Summary", {"A": 120})
    s.title("Risk Summary")
    s.section("RISK ANALYST — FINAL REVIEWED OUTPUT", 1)
    s.text(_version_line(result, RISK), italic=True, color=GREY)
    s.markdown(result.outputs[RISK], f"Source: final reviewed output — {RISK}")
    s.skip()
    s.section("IC PARTNER — FINAL QUALITY REVIEW", 1)
    source = "Source: final IC Partner review"
    s.text(f"Review round: {result.review_history[-1].round} of "
           f"{result.review_rounds_used} used", comment=source)
    s.text(f"Review status: {review.status}", bold=True)
    s.text(f"Summary: {review.summary}")
    for title, items in (("Cross-workstream issues", review.cross_workstream_issues),
                         ("Unresolved information gaps", review.unresolved_information_gaps)):
        s.skip()
        s.text(title, bold=True, color=NAVY, border=True)
        for item in items or ("None recorded in the final IC review.",):
            s.text(f"• {item}")
    s.print_setup()


def _build_memo(model: _Model, memo_result: MemoResult | None) -> None:
    s = _Sheet(model, "IC Memo", {"A": 120})
    s.title("Investment Committee Memo")
    if memo_result is None:
        s.text(NO_MEMO_NOTE, bold=True)
        return
    s.section("IC CONCLUSION", 1)
    s.text(memo_result.conclusion, bold=True, size=12, comment="Source: MemoResult.conclusion")
    s.skip()
    s.markdown(memo_result.memo_text, "Source: final IC memo (MemoResult.memo_text)")
    s.print_setup()


def _build_summary(model: _Model, deal: Deal, result: DealTeamResult, review: ICReview,
                   memo_result: MemoResult | None, valuation: dict) -> None:
    s = _Sheet(model, "IC Summary", {"A": 44, "B": 22, "C": 56})
    s.title("Investment Committee Summary")
    s.section("TRANSACTION")
    s.item("Deal", f"={s.ref('deal_name')}")
    s.item("Acquirer", f"={s.ref('acquirer.name')}")
    s.item("Target", f"={s.ref('target.name')}")
    s.item(_money_label("Offer enterprise value", deal),
           f"={s.ref('transaction_assumptions.offer_enterprise_value')}", fmt=FMT_MONEY)
    s.item(_money_label("Offer equity value", deal), f"={s.ref('tx.equity')}", fmt=FMT_MONEY)
    s.item("Cash %", f"={s.ref('tx.cash_pct')}", fmt=FMT_PCT)
    s.item("Stock %", f"={s.ref('tx.stock_pct')}", fmt=FMT_PCT)
    s.item("IC review status", result.status, comment="Source: DealTeamResult.status")
    s.item("Review rounds used", result.review_rounds_used, fmt=FMT_INT,
           comment="Source: DealTeamResult.review_rounds_used")
    s.skip()
    s.section(f"VALUATION SNAPSHOT ({deal.currency} {deal.units}, enterprise value)")
    summary = valuation["valuation_summary"]
    for method in summary["methods"]:
        label = f"{_short_method(method['method'])} EV"
        if method["applicable"]:
            s.item(label, f"={s.ref(METHOD_KEYS[method['method']])}", fmt=FMT_MONEY,
                   note=None if method["included_in_range"] else "Supplementary; not in range.")
        else:
            s.item(label, NOT_APPLICABLE, note=method["reason"])
    for label, key in (("Valuation low", "val.low"), ("Valuation high", "val.high")):
        if summary["range_available"]:
            s.item(label, f"={s.ref(key)}", fmt=FMT_MONEY)
        else:
            s.item(label, "Not available", note=summary["range_note"])
    s.item("Offer EV", f"={s.ref('val.offer_ev')}", fmt=FMT_MONEY)
    s.skip()
    s.section(f"SYNERGY SNAPSHOT ({deal.currency} {deal.units})")
    s.item("PV of synergies", f"={s.ref('syn.pv')}", fmt=FMT_MONEY)
    s.item("Integration cost", f"={s.ref('syn.integration_cost')}", fmt=FMT_MONEY)
    s.item("Net synergy value", f"={s.ref('syn.net')}", fmt=FMT_MONEY,
           note="Based on unvalidated synergy assumptions.")
    s.skip()
    s.section("TRANSACTION SNAPSHOT")
    s.item(_money_label("Cash consideration", deal), f"={s.ref('tx.cash_consideration')}",
           fmt=FMT_MONEY, note="Debt funded under the model convention; not a confirmed plan.")
    s.item(_money_label("Stock consideration", deal), f"={s.ref('tx.stock_consideration')}",
           fmt=FMT_MONEY)
    s.item(f"New shares issued ({deal.units})", f"={s.ref('tx.new_shares')}", fmt=FMT_SHARES)
    s.item("Seller ownership", f"={s.ref('tx.seller_ownership')}", fmt=FMT_OWNERSHIP)
    s.item("EPS result — excluding synergies", f"={s.ref('acc.result.A')}")
    s.item("EPS result — including run-rate synergies", f"={s.ref('acc.result.B')}",
           note="Model scenario only; not the expected outcome.")
    s.skip()
    s.section("IC CONCLUSION")
    if memo_result is None:
        s.text("No IC memo supplied", bold=True)
    else:
        s.text(memo_result.conclusion, bold=True, size=12, comment="Source: MemoResult.conclusion")
    s.skip()
    s.section("KEY UNRESOLVED ITEMS (FINAL IC REVIEW)")
    for gap in review.unresolved_information_gaps or ("None recorded in the final IC review.",):
        s.text(f"• {gap}", comment="Source: final IC Partner review — unresolved_information_gaps")
    s.freeze("A4")
    s.print_setup()


# --- Validation and export ---------------------------------------------------------------------

def _close(actual: Any, expected: Any) -> bool:
    if isinstance(expected, str) or isinstance(actual, str):
        return actual == expected
    return math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-9)


def validate_workbook(model: _Model) -> None:
    """Evaluate every formula and compare model cells against deterministic Python."""
    for ws in model.workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and ("[" in cell.value or "#REF!" in cell.value):
                    raise ExcelExportError(
                        f"{ws.title}!{cell.coordinate} has an external or broken reference")
    evaluator = WorkbookEvaluator(model.workbook)
    try:
        evaluator.evaluate_all()
    except FormulaError as exc:
        raise ExcelExportError(f"workbook formula error: {exc}") from exc
    mismatches = [
        f"{sheet}!{coordinate}: workbook {evaluator.value(sheet, coordinate)!r} != "
        f"Python {expected!r}"
        for sheet, coordinate, expected in model.expectations
        if not _close(evaluator.value(sheet, coordinate), expected)
    ]
    if mismatches:
        raise ExcelExportError("workbook disagrees with deterministic Python: "
                               + "; ".join(mismatches))


def build_workbook(
    deal: Deal, result: DealTeamResult, memo_result: MemoResult | None = None,
    *, deal_source: str = "data/sample_deal.json",
    input_sources: Mapping[str, str] | None = None,
) -> Workbook:
    """Build and validate the workbook in memory (does not save)."""
    review = check_export_eligibility(deal, result, memo_result)
    # Full-precision deterministic outputs for formula checks (display rounding off).
    valuation = build_valuation_context(deal, ndigits=None)
    synergy = build_synergy_context(deal, ndigits=None)
    transaction = build_transaction_structure_context(deal, ndigits=None)

    workbook = Workbook()
    # openpyxl writes formulas without cached results; tell Excel to calculate
    # everything itself when the file is opened.
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.remove(workbook.active)
    for title in SHEET_ORDER:
        workbook.create_sheet(title)
    model = _Model(workbook)

    # input lineage recorded in the approved result (share-basis definition, data quality)
    _, data_quality = lineage_from_contexts(result.deterministic_contexts)
    _build_assumptions(model, deal, deal_source, input_sources,
                       result.deterministic_contexts["transaction_structure"]["calculation_notes"]
                       .get("acquirer_share_basis_definition"), data_quality)
    _build_transaction(model, deal, transaction)
    _build_synergies(model, deal, synergy, synergy["labels"])
    _build_valuation(model, deal, valuation, valuation["transaction"])
    _build_accretion(model, deal, transaction)
    _build_due_diligence(model, result)
    _build_risk(model, result, review)
    _build_memo(model, memo_result)
    _build_summary(model, deal, result, review, memo_result, valuation)

    validate_workbook(model)
    return workbook


def export_ma_workbook(
    deal: Deal, result: DealTeamResult, memo_result: MemoResult | None,
    output_path: str | Path = DEFAULT_OUTPUT, *, overwrite: bool = False,
    deal_source: str = "data/sample_deal.json",
    input_sources: Mapping[str, str] | None = None,
) -> Path:
    """Build, validate, and save the workbook. Refuses to overwrite an existing file
    unless overwrite=True."""
    path = Path(output_path)
    if path.suffix.lower() != ".xlsx":
        raise ExcelExportError("output path must end in .xlsx")
    if path.exists() and not overwrite:
        raise ExcelExportError(f"{path} already exists; pass overwrite=True to replace it")
    workbook = build_workbook(deal, result, memo_result, deal_source=deal_source,
                              input_sources=input_sources)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
