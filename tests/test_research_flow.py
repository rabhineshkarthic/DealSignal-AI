"""Research flow: registry, draft, approval/overrides, Deal construction, Excel
provenance, end-to-end mocked pipeline, and the Streamlit research UI.

No SEC request, no market-price request, no Gemini request, no network.
"""

import io
import socket
from pathlib import Path

import pytest
from openpyxl import load_workbook

import llm_client
import ui_helpers as ui
from deal_input import DealInputError, form_from_deal
from memo import InvestmentCommitteeMemoGenerator
from models import Deal, load_deal
from orchestrator import APPROVED, DealTeamOrchestrator
from research import registry
from research.draft import create_draft
from research.errors import ResearchError
from research.market_data import StooqMarketPriceProvider, missing_price, quote_to_value
from research.models import ResearchStatus
from research.provenance import ASSUMPTION_PATHS, approve_draft, build_researched_deal
from research_fixtures import FakeMarketFetch, FakeSECFetch, sec_provider
from test_memo import ScriptedLLM, approved, memo_json

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = load_deal(ROOT / "data" / "sample_deal.json")
US = registry.get_jurisdiction("US")


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def blocked(*args, **kwargs):
        raise AssertionError("real network access attempted")

    def no_client(*args, **kwargs):
        raise AssertionError("a real GeminiClient must not be created")

    monkeypatch.setattr(socket.socket, "connect", blocked)
    monkeypatch.setattr("urllib.request.urlopen", blocked)
    monkeypatch.setattr(llm_client.GeminiClient, "__init__", no_client)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)


def researched(price_body=None):
    provider = sec_provider()
    acquirer = provider.research_company(provider.resolve_company("ALPH"))
    target = provider.research_company(provider.resolve_company("BETA"))
    if price_body is None:
        price = quote_to_value(StooqMarketPriceProvider(fetch=FakeMarketFetch()).get_quote("ALPH", "US"))
    else:
        price = missing_price(price_body)
    return create_draft(US, acquirer, target, price)


def assumptions(offer_ev=60000.0):
    values = {p: v for p, v in form_from_deal(SAMPLE).items() if p in ASSUMPTION_PATHS}
    values["transaction_assumptions.offer_enterprise_value"] = offer_ev
    return values


# Beta's EBITDA and FCF are MISSING in the fixtures, so they must be entered.
FILL_MISSING = {"target.ebitda": 8900.0, "target.free_cash_flow": 7000.0}


# --- Registry --------------------------------------------------------------------------------

def test_registry_us_is_the_only_available_jurisdiction():
    assert [j.label for j in registry.available_jurisdictions()] == ["United States — SEC"]
    assert [j.label for j in registry.planned_jurisdictions()] == [
        "France — coming later", "Germany — coming later"]


@pytest.mark.parametrize("code", ["FR", "DE"])
def test_planned_jurisdictions_cannot_create_providers(code):
    with pytest.raises(ResearchError, match="not supported yet"):
        registry.create_provider(code)
    assert registry.create_market_price_provider(code) is None


def test_unknown_jurisdiction_rejected():
    with pytest.raises(ResearchError, match="Unknown"):
        registry.get_jurisdiction("XX")


# --- Draft ---------------------------------------------------------------------------------------

def test_draft_rows_and_missing_fields():
    draft = researched()
    paths = [r.path for r in draft.rows()]
    assert paths == ["acquirer.revenue", "acquirer.ebitda", "acquirer.net_income", "acquirer.cash",
                     "acquirer.debt", "acquirer.shares_outstanding", "acquirer.share_price",
                     "target.revenue", "target.ebitda", "target.net_income",
                     "target.free_cash_flow", "target.cash", "target.debt"]
    assert draft.missing_paths() == ("target.ebitda", "target.free_cash_flow")
    assert draft.share_price.status == ResearchStatus.MARKET


def test_market_price_failure_is_non_fatal():
    draft = researched(price_body="Market price unavailable.")
    assert draft.share_price.status == ResearchStatus.MISSING
    assert "acquirer.share_price" in draft.missing_paths()
    assert draft.acquirer.fields["revenue"].available  # research continued


def test_currency_mismatch_rejected():
    from dataclasses import replace
    draft = researched()
    with pytest.raises(ResearchError, match="different currencies"):
        create_draft(US, draft.acquirer, replace(draft.target, reporting_currency="EUR"),
                     draft.share_price)


# --- Approval and overrides ------------------------------------------------------------------------

def test_missing_values_must_be_entered():
    with pytest.raises(ResearchError) as exc_info:
        approve_draft(researched())
    message = str(exc_info.value)
    assert "BETA ANALYTICS INC — EBITDA is missing" in message
    assert "Enter EBITDA manually before building the deal" in message
    assert "Free cash flow is missing" in message


def test_accepting_researched_values_keeps_their_status():
    provenance = approve_draft(researched(), FILL_MISSING)
    revenue = provenance.approved["acquirer.revenue"]
    assert (revenue.value, revenue.status) == (245122.0, ResearchStatus.REPORTED)
    assert provenance.approved["acquirer.ebitda"].status == ResearchStatus.DERIVED
    assert provenance.approved["acquirer.share_price"].status == ResearchStatus.MARKET


def test_override_keeps_the_original_research_value():
    provenance = approve_draft(researched(), {**FILL_MISSING, "acquirer.revenue": 250000.0},
                               {"acquirer.revenue": "Pro forma for a closed acquisition"})
    override = provenance.approved["acquirer.revenue"]
    assert override.status == ResearchStatus.USER_OVERRIDE and override.value == 250000.0
    assert override.original.status == ResearchStatus.REPORTED
    assert override.original.normalized_value == 245122.0
    assert override.original.concept == "RevenueFromContractWithCustomerExcludingAssessedTax"
    assert override.override_reason == "Pro forma for a closed acquisition"


def test_entering_a_missing_value_is_a_user_override_of_missing():
    provenance = approve_draft(researched(), FILL_MISSING)
    ebitda = provenance.approved["target.ebitda"]
    assert ebitda.status == ResearchStatus.USER_OVERRIDE and not ebitda.original.available
    assert len(provenance.overrides) == 2


def test_unchanged_value_is_not_an_override():
    provenance = approve_draft(researched(), {**FILL_MISSING, "acquirer.revenue": 245122.0})
    assert provenance.approved["acquirer.revenue"].status == ResearchStatus.REPORTED


def test_override_must_be_a_finite_number():
    with pytest.raises(ResearchError, match="finite number"):
        approve_draft(researched(), {**FILL_MISSING, "acquirer.debt": float("nan")})


# --- Building the existing Deal ----------------------------------------------------------------------

def test_researched_deal_is_the_existing_deal_type():
    draft = researched()
    provenance = approve_draft(draft, FILL_MISSING)
    deal = build_researched_deal(draft, provenance, assumptions())
    assert type(deal) is Deal
    assert deal.acquirer.name == "ALPHA SOFTWARE CORP" and deal.target.name == "BETA ANALYTICS INC"
    assert deal.deal_name == "ALPHA SOFTWARE CORP acquiring BETA ANALYTICS INC"
    assert (deal.currency, deal.units) == ("USD", "millions")
    assert deal.acquirer.revenue == 245122.0 and deal.acquirer.share_price == 412.37
    assert deal.acquirer.shares_outstanding == 7433.038381
    assert deal.target.ebitda == 8900.0 and deal.target.debt == 9429.0
    assert deal.dcf_assumptions.wacc == SAMPLE.dcf_assumptions.wacc  # user assumption


def test_deterministic_contexts_for_researched_deal():
    draft = researched()
    deal = build_researched_deal(draft, approve_draft(draft, FILL_MISSING), assumptions())
    contexts = DealTeamOrchestrator.build_deterministic_contexts(deal)
    assert contexts["valuation"]["comparable_company"]["enterprise_value"] == pytest.approx(8900.0 * 10.5)
    equity = contexts["transaction_structure"]["calculation"]["offer_equity_value"]
    assert equity == pytest.approx(60000.0 - (9429.0 - 8848.0))
    assert dict(ui.deal_preview(deal, contexts))["Offer EV"] == "USD 60.0bn"  # display only
    assert deal.transaction_assumptions.offer_enterprise_value == 60000.0   # stored in millions


def test_assumptions_are_required_and_validated():
    draft = researched()
    provenance = approve_draft(draft, FILL_MISSING)
    values = assumptions()
    del values["transaction_assumptions.offer_enterprise_value"]
    with pytest.raises(DealInputError, match="Missing fields"):
        build_researched_deal(draft, provenance, values)
    with pytest.raises(DealInputError, match="exceeds enterprise value"):
        build_researched_deal(draft, provenance, assumptions(offer_ev=100.0))


# --- Excel provenance ---------------------------------------------------------------------------------

def test_excel_sources_describe_research_provenance():
    draft = researched()
    provenance = approve_draft(draft, {**FILL_MISSING, "acquirer.revenue": 250000.0})
    sources = provenance.excel_sources()
    assert sources["acquirer.revenue"].startswith("EFFECTIVE VALUE 250000.0 — user override")
    assert ("Original researched value (kept for provenance, not used): 245122.0 USD millions "
            "— REPORTED — SEC EDGAR") in sources["acquirer.revenue"]
    assert sources["target.revenue"] == (
        "SEC EDGAR — BETA ANALYTICS INC — Form 10-K/A — us-gaap:Revenues — 2024-01-01 to "
        "2024-12-31 — document 0000000202-25-000009")
    assert sources["acquirer.ebitda"].startswith("SEC EDGAR — ALPHA SOFTWARE CORP — Form 10-K — "
                                                 "derived: EBITDA = operating income")
    assert sources["acquirer.share_price"].startswith("Stooq (third-party market data) — as of")
    assert sources["acquirer.name"] == "SEC EDGAR — registrant name (CIK 0000000101)"
    assert sources["dcf_assumptions.wacc"] == "User assumption (not researched)"
    assert "custom deal input" not in " ".join(sources.values())


# --- End-to-end mocked flow -------------------------------------------------------------------------------

def test_end_to_end_research_to_memo_and_excel():
    fetch = FakeSECFetch()
    provider = sec_provider(fetch)
    market_fetch = FakeMarketFetch()
    acquirer = provider.research_company(provider.resolve_company("ALPH"))
    target = provider.research_company(provider.resolve_company("BETA"))
    price = quote_to_value(StooqMarketPriceProvider(fetch=market_fetch).get_quote("ALPH", "US"))
    draft = create_draft(US, acquirer, target, price)
    provenance = approve_draft(draft, FILL_MISSING)
    deal = build_researched_deal(draft, provenance, assumptions())

    llm = ScriptedLLM([approved()], memo_response=memo_json())
    result = DealTeamOrchestrator(llm).run(deal)
    memo = InvestmentCommitteeMemoGenerator(llm).generate(result, deal)
    data = ui.workbook_bytes(deal, result, memo, ui.CUSTOM_SOURCE_LABEL, provenance.excel_sources())

    assert result.status == APPROVED and memo.deal_name == deal.deal_name
    wb = load_workbook(io.BytesIO(data))
    comments = {c.comment.text for row in wb["Assumptions"].iter_rows() for c in row if c.comment}
    assert any(t.startswith("Source: SEC EDGAR — ALPHA SOFTWARE CORP — Form 10-K — us-gaap:"
                            "RevenueFromContractWithCustomerExcludingAssessedTax") for t in comments)
    assert "Source: User assumption (not researched)" in comments
    assert len(fetch.calls) == 5 and len(market_fetch.calls) == 1  # all fake, none real
    assert ui.workbook_filename(deal) == "ALPHA_BETA_MA_Model.xlsx"


# --- Streamlit research UI (AppTest) -------------------------------------------------------------------------

@pytest.fixture
def fake_providers(monkeypatch):
    monkeypatch.setitem(registry.PROVIDER_FACTORIES, "sec", lambda: sec_provider())
    monkeypatch.setitem(registry.MARKET_PRICE_FACTORIES, "US",
                        lambda: StooqMarketPriceProvider(fetch=FakeMarketFetch()))
    from config import load_settings
    monkeypatch.setattr(ui, "load_settings", lambda: load_settings(env_file=None))


def app_test():
    from streamlit.testing.v1 import AppTest
    return AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)


def open_research(at, acquirer="ALPH", target="BETA"):
    at.sidebar.radio(key="deal_source").set_value(ui.RESEARCH_DEAL).run()
    at.text_input(key="research_q_acquirer").set_value(acquirer)
    at.text_input(key="research_q_target").set_value(target)
    next(b for b in at.button if b.label == "Research Companies").click().run()


def text(at):
    return " ".join(m.value for m in at.markdown)


def test_app_research_flow_builds_deal(fake_providers):
    at = app_test().run()
    open_research(at)
    assert not at.exception
    page = text(at)
    for fragment in ["Reported", "Derived", "Market data", "Missing", "as of 2024-09-30",
                     "RevenueFromContractWithCustomerExcludingAssessedTax", "France — coming later"]:
        assert fragment in page, fragment
    assert [s.options for s in at.selectbox] == [
        ["United States — SEC"],
        # share-basis definition for an acquirer share override (never relabelled)
        ["Point-in-time shares outstanding", "Weighted-average basic shares (modelling proxy)",
         "Weighted-average diluted shares (modelling proxy)"]]
    labels = [e.label for e in at.expander]
    assert "Original research record — ALPHA SOFTWARE CORP · Revenue" in labels

    at.number_input(key="rv_target.ebitda").set_value(8900.0)
    at.number_input(key="rv_target.free_cash_flow").set_value(7000.0)
    at.number_input(key="ra_transaction_assumptions.offer_enterprise_value").set_value(60000.0)
    at.button(key="FormSubmitter:research_review_form-Build Deal").click().run()
    assert not at.exception
    deal = at.session_state["researched_deal"]
    assert deal.acquirer.revenue == 245122.0 and deal.target.ebitda == 8900.0
    provenance = at.session_state["research_provenance"]
    assert provenance.approved["acquirer.revenue"].status == ResearchStatus.REPORTED  # untouched
    assert {o.path for o in provenance.overrides} == {"target.ebitda", "target.free_cash_flow"}
    assert "Deal created · ALPHA SOFTWARE CORP → BETA ANALYTICS INC" in text(at)
    assert "Researched public companies" in text(at)


def test_app_missing_values_block_deal_creation(fake_providers):
    at = app_test().run()
    open_research(at)
    at.number_input(key="ra_transaction_assumptions.offer_enterprise_value").set_value(60000.0)
    at.button(key="FormSubmitter:research_review_form-Build Deal").click().run()
    assert at.session_state["researched_deal"] is None
    assert any("Enter EBITDA manually before building the deal" in e.value for e in at.error)


def test_app_ambiguous_company_requires_selection(fake_providers):
    at = app_test().run()
    open_research(at, acquirer="ACME")
    assert not at.exception
    assert at.session_state["research_draft"] is None
    radio = next(r for r in at.radio if r.label.startswith("Several companies match the acquirer"))
    assert list(radio.options) == [
        "ACME CLOUD CORP · Ticker: ACMC · CIK: 0000000303",
        "ACME HOLDINGS INC · Ticker: ACMH · CIK: 0000000304"]


def test_app_new_research_clears_stale_analysis(fake_providers):
    at = app_test().run()
    open_research(at)
    at.number_input(key="rv_target.ebitda").set_value(8900.0)
    at.number_input(key="rv_target.free_cash_flow").set_value(7000.0)
    at.number_input(key="ra_transaction_assumptions.offer_enterprise_value").set_value(60000.0)
    at.button(key="FormSubmitter:research_review_form-Build Deal").click().run()
    stale = DealTeamOrchestrator(ScriptedLLM([approved()])).run(at.session_state["researched_deal"])
    at.session_state["deal_team_result"] = stale
    at.session_state["memo_result"] = "stale memo"
    next(b for b in at.button if b.label == "Research Companies").click().run()
    assert not at.exception
    assert at.session_state["researched_deal"] is None
    assert at.session_state["deal_team_result"] is None and at.session_state["memo_result"] is None


def test_app_sample_and_custom_modes_unaffected(fake_providers):
    at = app_test().run()
    assert not at.exception and at.session_state["deal_source"] == ui.SAMPLE_DEAL
    at.sidebar.radio(key="deal_source").set_value(ui.CUSTOM_DEAL).run()
    assert not at.exception
    assert [r.value for r in at.sidebar.radio] == [ui.OFFLINE_DEMO, ui.CUSTOM_DEAL]


def test_clear_research_state_drops_stale_form_values():
    state = {"researched_deal": "D", "research_draft": "X", "rv_acquirer.revenue": 1.0,
             "ra_dcf_assumptions.wacc": 9.0, "rr_target.debt": "why", "mode": "keep"}
    ui.clear_research_state(state)
    assert state["researched_deal"] is None and state["research_draft"] is None
    assert not any(k.startswith(("rv_", "ra_", "rr_")) for k in state)
    assert state["mode"] == "keep"
