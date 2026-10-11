"""Phase 6: Streamlit app helpers, offline demo state, and an AppTest smoke run.

No browser, no network, no Gemini, no API key. The real GeminiClient is booby-trapped
in every test so any attempt to create one fails loudly.
"""

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

import llm_client
import ui_helpers as ui
from agent_validation import AgentValidationError, NumericViolation
from agents.registry import AGENT_NAMES, FINANCIAL_DD, RISK
from config import ConfigError, GeminiSettings, load_settings
from demo_state import DEMO_CONCLUSION, DEMO_LABEL, OfflineDemoLLM, build_demo_state
from excel_export import SHEET_ORDER, ExcelExportError
from ic_review import ICReviewParseError
from memo import MemoGenerationError, MemoParseError
from orchestrator import APPROVED, DealTeamOrchestrator, OrchestrationError
from test_memo import ScriptedLLM, revise

ROOT = Path(__file__).resolve().parent.parent
SECRET = "test-secret-key-should-never-appear"


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    def forbidden(*args, **kwargs):
        raise AssertionError("a real GeminiClient must not be created offline")

    monkeypatch.setattr(llm_client.GeminiClient, "__init__", forbidden)


@pytest.fixture(scope="module")
def demo():
    return build_demo_state()


# --- Demo state -------------------------------------------------------------------------------

def test_demo_state_builds_without_api_key_or_client(demo):
    assert demo.is_demo and demo.label == DEMO_LABEL
    assert demo.result.status == APPROVED
    assert set(demo.deterministic_contexts) == {"valuation", "synergy", "transaction_structure"}


def test_demo_content_is_clearly_labelled(demo):
    for name in AGENT_NAMES:
        assert DEMO_LABEL in demo.result.outputs[name], name
    assert DEMO_LABEL in demo.memo.memo_text
    assert demo.memo.conclusion == DEMO_CONCLUSION


def test_demo_runs_through_real_controls_and_records_a_revision(demo):
    (record,) = demo.result.revision_history
    assert record.directly_flagged == (FINANCIAL_DD,)
    assert ui.revision_flows(demo.result) == [[
        "Financial Due Diligence revised", "Deal Structuring Analyst stale rerun",
        "Risk Analyst stale rerun", "IC re-review (round 2)"]]
    revised = demo.result.outputs[FINANCIAL_DD]
    assert "cannot yet be assessed" in revised
    assert "The target has customer concentration risk." not in revised


def test_demo_deterministic_values_are_real(demo):
    from deal_analysis import build_valuation_context
    assert demo.deterministic_contexts["valuation"] == build_valuation_context(demo.deal)


def test_demo_client_never_uses_network():
    llm = OfflineDemoLLM()
    assert llm.last_model_used == "offline-demo (no LLM)"
    assert not isinstance(llm, llm_client.GeminiClient)


# --- Formatting and view data ------------------------------------------------------------------------

def test_formatting(demo):
    d = demo.deal
    assert ui.fmt_money(205.0, d) == "USD 205.0m"
    assert ui.fmt_money(-7.5, d) == "(USD 7.5m)"
    assert ui.fmt_money(172.024, d, 4) == "USD 172.0240m"
    assert ui.fmt_pct(0.055) == "5.5%"
    assert ui.fmt_multiple(10.5) == "10.5x"
    assert ui.fmt_eps(0.7325) == "0.7325"
    assert ui.fmt_shares(2.96, d) == "2.96m shares"


def test_overview_uses_deterministic_values(demo):
    metrics = dict(ui.overview_metrics(demo.deal, demo.deterministic_contexts))
    assert metrics == {"Offer Enterprise Value": "USD 205.0m", "Offer Equity Value": "USD 185.0m",
                       "Cash Consideration": "USD 111.0m", "Stock Consideration": "USD 74.0m",
                       "Target Net Debt": "USD 20.0m"}
    acquirer, target = ui.company_rows(demo.deal)
    assert dict(acquirer)["Revenue"] == "USD 650.0m"
    assert dict(target)["Free Cash Flow"] == "USD 10.0m"


def test_valuation_view(demo):
    view = ui.valuation_view(demo.deal, demo.deterministic_contexts)
    assert view["chart"] == {"Comparable Companies": 189.0, "Precedent Transactions": 216.0,
                             "Discounted Cash Flow": 172.024, "Offer EV": 205.0}
    assert view["offer_vs_range"] == "within"
    assert [r["Year"] for r in view["schedule"]] == [1, 2, 3, 4, 5]


def test_synergy_and_transaction_views(demo):
    syn = ui.synergy_view(demo.deal, demo.deterministic_contexts)
    assert dict(syn["metrics"])["Net Synergy Value"] == "USD 13.5878m"
    assert len(syn["schedule"]) == 5
    tx = ui.transaction_view(demo.deal, demo.deterministic_contexts)
    assert [c["title"] for c in tx["cases"]] == ["A. Excluding synergies",
                                                 "B. Including full run-rate synergies"]
    assert [c["result"] for c in tx["cases"]] == ["Accretive", "Accretive"]
    assert dict(tx["shares"])["New Shares Issued"] == "2.96m shares"


def test_agent_cards_show_current_versions(demo):
    cards = {c.name: c for c in ui.agent_cards(demo.result)}
    assert list(cards) == list(AGENT_NAMES)
    assert (cards[FINANCIAL_DD].version, cards[FINANCIAL_DD].kind) == (2, "direct_revision")
    assert (cards[RISK].version, cards[RISK].kind) == (2, "stale_rerun")
    assert cards[FINANCIAL_DD].output == demo.result.outputs[FINANCIAL_DD]
    assert all(c.output is None for c in ui.agent_cards(None))


def test_review_rounds_view(demo):
    rounds = ui.review_rounds(demo.result)
    assert [r["status"] for r in rounds] == ["revision_required", "approved"]
    assert rounds[0]["revision_requests"][0][0] == FINANCIAL_DD
    assert ui.review_rounds(None) == [] and ui.revision_flows(None) == []


def test_no_revisions_means_no_flow(deal=None):
    from test_memo import approved
    from models import load_deal
    result = DealTeamOrchestrator(ScriptedLLM([approved()])).run(
        load_deal(ROOT / "data" / "sample_deal.json"))
    assert ui.revision_flows(result) == []


# --- Action availability ---------------------------------------------------------------------------

def test_action_state(demo):
    offline = ui.action_state(ui.OFFLINE_DEMO, False, demo.result)
    assert not offline.can_run and not offline.can_generate_memo and offline.can_export
    live_unconfigured = ui.action_state(ui.LIVE_ANALYSIS, False, None)
    assert not live_unconfigured.can_run and not live_unconfigured.can_export
    live_ready = ui.action_state(ui.LIVE_ANALYSIS, True, None)
    assert live_ready.can_run and not live_ready.can_generate_memo
    assert live_ready.memo_block_reason == "Run the full deal team first."
    approved = ui.action_state(ui.LIVE_ANALYSIS, True, demo.result)
    assert approved.can_generate_memo and approved.can_export


def test_unresolved_review_blocks_memo_and_export():
    from models import load_deal
    result = DealTeamOrchestrator(ScriptedLLM([revise(RISK), revise(RISK)])).run(
        load_deal(ROOT / "data" / "sample_deal.json"))
    state = ui.action_state(ui.LIVE_ANALYSIS, True, result)
    assert not state.can_generate_memo and not state.can_export
    assert "review_limit_reached" in state.memo_block_reason


# --- Settings and errors -----------------------------------------------------------------------------

def test_live_settings_summary_without_key(monkeypatch):
    monkeypatch.setattr(ui, "load_settings", lambda: load_settings(env_file=None))
    summary = ui.live_settings_summary()
    assert not summary.configured and "GEMINI_API_KEY is not set" in summary.message


def test_live_settings_summary_never_exposes_key(monkeypatch):
    monkeypatch.setattr(ui, "load_settings",
                        lambda: GeminiSettings(api_key=SECRET, model="m", fallback_models=("f",)))
    summary = ui.live_settings_summary()
    assert summary.configured and summary.primary_model == "m"
    assert SECRET not in repr(summary)


def _orch_error(cause, stage="initial_analysis"):
    from orchestrator import DealTeamResult
    return OrchestrationError(stage, cause, DealTeamResult(outputs={}, deterministic_contexts={}),
                              agent="Deal Lead")


@pytest.mark.parametrize("exc, title", [
    (ConfigError("GEMINI_API_KEY is not set."), "Gemini is not configured"),
    (_orch_error(llm_client.LLMError("Gemini daily quota exhausted for model 'm'")), "LLM error"),
    (_orch_error(llm_client.LLMError("HTTP 503 after 4 attempts")), "LLM error"),
    (_orch_error(AgentValidationError("Deal Lead", [NumericViolation("25%", 25, 1, "25%")], "x")),
     "Output rejected by the numeric provenance validator"),
    (_orch_error(ICReviewParseError("bad"), stage="ic_review_parse"), "IC review rejected"),
    (MemoGenerationError("unresolved"), "Memo generation blocked"),
    (MemoParseError("bad"), "Memo rejected"),
    (ExcelExportError("bad"), "Excel export failed"),
])
def test_describe_error(exc, title):
    message = ui.describe_error(exc)
    assert message.title.startswith(title)
    assert SECRET not in message.details


def test_quota_and_capacity_messages():
    quota = ui.describe_error(_orch_error(llm_client.LLMError("daily quota exhausted")))
    capacity = ui.describe_error(_orch_error(llm_client.LLMError("HTTP 503 after 4 attempts")))
    assert "quota" in quota.message and "capacity" in capacity.message
    assert "during initial analysis (Deal Lead)" in quota.title


# --- Downloads ------------------------------------------------------------------------------------------

def test_workbook_bytes_from_demo_state(demo):
    data = ui.workbook_bytes(demo.deal, demo.result, demo.memo)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == list(SHEET_ORDER)
    dd = [c.value for row in wb["Due Diligence"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(DEMO_LABEL in t for t in dd)
    assert ui.workbook_filename(demo.deal) == "Northstar_Atlas_MA_Model.xlsx"


def test_memo_download_is_exact_memo_text(demo):
    assert ui.memo_download_bytes(demo.memo) == demo.memo.memo_text.encode("utf-8")


def test_event_labels(demo):
    runs = demo.result.agent_runs
    assert ui.event_label(runs[0]) == "[1] Deal Lead completed"
    revised = next(r for r in runs if r.kind == "direct_revision")
    assert ui.event_label(revised) == "[5A] Financial Due Diligence (revised) completed"
    assert ui.event_label(demo.result.review_history[1]) == "IC Partner review round 2: approved"


# --- Streamlit AppTest --------------------------------------------------------------------------------

def app_test():
    from streamlit.testing.v1 import AppTest
    return AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)


@pytest.fixture
def unconfigured(monkeypatch):
    monkeypatch.setattr(ui, "load_settings", lambda: load_settings(env_file=None))


def page_text(at):
    return " ".join(m.value for m in at.markdown)


def test_app_offline_demo_renders(unconfigured):
    at = app_test().run()
    assert not at.exception
    assert at.session_state["mode"] == ui.OFFLINE_DEMO  # offline is the default
    text = page_text(at)
    for value in ["USD 205.0m", "USD 185.0m", "USD 111.0m", "USD 74.0m", "USD 20.0m",
                  "USD 172.0m", "USD 189.0m", "USD 216.0m"]:
        assert value in text, value
    assert "Offline demo" in text and "Illustrative qualitative analysis" in text
    assert DEMO_LABEL in text
    assert ui.SYNERGY_WARNING in text and ui.SYNERGY_CASE_CAUTION in text
    assert "Financing modelling convention" in text
    assert DEMO_CONCLUSION in text


def test_app_has_five_primary_sections_with_sub_tabs(unconfigured):
    at = app_test().run()
    labels = [t.label for t in at.tabs]
    assert labels == ["Overview", "Financial Analysis", "Valuation", "Synergies",
                      "Transaction & Accretion", "Deal Team", "IC Review", "Deliverables",
                      "IC Memo", "Excel Model", "Architecture"]
    import app
    assert app.PRIMARY_SECTIONS == ("Overview", "Financial Analysis", "Deal Team",
                                    "IC Review", "Deliverables")


def test_app_ic_review_renders_revision_history(unconfigured):
    text = page_text(app_test().run())
    assert "ROUND 1" in text and "ROUND 2" in text
    for step in ["Financial Due Diligence", "direct revision", "Deal Structuring Analyst",
                 "stale rerun", "Risk Analyst"]:
        assert step in text
    assert "Customer, market, and product data are not supplied." in text


def test_app_deal_team_expanders_show_current_versions(unconfigured):
    at = app_test().run()
    labels = [e.label for e in at.expander]
    assert "5A · Financial Due Diligence — Version 2 · Direct revision" in labels
    assert "7 · Risk Analyst — Version 2 · Stale rerun" in labels
    assert "1 · Deal Lead — Version 1 · Initial" in labels


def test_app_offline_generate_workbook_and_memo_download(unconfigured):
    at = app_test().run()
    next(b for b in at.button if b.label == "Generate Excel Model").click().run()
    assert not at.exception
    assert at.session_state["excel_bytes"][:2] == b"PK"  # a real .xlsx (zip) was produced


def test_app_live_mode_without_key_disables_actions(unconfigured):
    at = app_test().run()
    at.sidebar.radio[0].set_value(ui.LIVE_ANALYSIS).run()
    assert not at.exception
    assert any("GEMINI_API_KEY is not set" in w.value for w in at.sidebar.warning)
    buttons = {b.label: b for b in at.sidebar.button}
    assert buttons["Run Full Deal Team"].disabled
    assert buttons["Generate IC Memo"].disabled


def test_app_live_flow_with_scripted_client(monkeypatch):
    """The live buttons, wired to the real orchestrator and memo generator, with the
    scripted demo client standing in for Gemini."""
    monkeypatch.setattr(ui, "live_settings_summary", lambda: ui.LiveSettingsSummary(
        True, "primary-model", ("fallback-model",), None,
        GeminiSettings(api_key=SECRET, model="primary-model")))
    monkeypatch.setattr(ui, "create_live_client", lambda settings: OfflineDemoLLM())
    at = app_test().run()
    at.sidebar.radio[0].set_value(ui.LIVE_ANALYSIS).run()
    sidebar_text = " ".join(m.value for m in at.sidebar.markdown)
    assert "primary-model" in sidebar_text and SECRET not in sidebar_text
    assert "Advanced settings" in [e.label for e in at.sidebar.expander]
    next(b for b in at.sidebar.button if b.label == "Run Full Deal Team").click().run()
    assert not at.exception
    result = at.session_state["deal_team_result"]
    assert result.status == APPROVED and result.review_rounds_used == 2
    next(b for b in at.sidebar.button if b.label == "Generate IC Memo").click().run()
    assert not at.exception
    assert at.session_state["memo_result"].conclusion == DEMO_CONCLUSION


# --- Visual components (pure HTML builders) ----------------------------------------------------------

def test_components_escape_dynamic_text():
    import ui_components as uc
    hostile = "<script>alert(1)</script>"
    for markup in [uc.kpi_cards([(hostile, hostile, hostile)]), uc.kv_rows([(hostile, hostile)]),
                   uc.callout(hostile, "amber", hostile), uc.conclusion(hostile, hostile),
                   uc.bullet_list([hostile]), uc.flow([[hostile]]), uc.sheet_grid([hostile]),
                   uc.bar_comparison([(hostile, 1.0, hostile, False)]),
                   uc.timeline([{"round": 1, "status": "approved", "summary": hostile,
                                 "requests": [(hostile, hostile)], "reruns": [(hostile, hostile)]}])]:
        assert "<script>" not in markup and "&lt;script&gt;" in markup


def test_bar_comparison_scales_for_display_only():
    import ui_components as uc
    markup = uc.bar_comparison([("A", 50.0, "50", False), ("B", 100.0, "100", True)])
    assert "width:50.0%" in markup and "width:100.0%" in markup
    assert markup.count("highlight") == 1


def test_review_timeline_uses_audit_trail(demo):
    rounds = ui.review_timeline(demo.result)
    assert [r["status"] for r in rounds] == ["revision_required", "approved"]
    assert rounds[0]["reruns"] == [(FINANCIAL_DD, "direct revision"),
                                   ("Deal Structuring Analyst", "stale rerun"),
                                   (RISK, "stale rerun")]
    assert rounds[1]["reruns"] == []
    assert rounds[0]["requests"][0][0] == FINANCIAL_DD
    assert ui.review_timeline(None) == []


def test_kind_labels():
    assert [ui.kind_label(k) for k in ("initial", "direct_revision", "stale_rerun", None)] == [
        "Initial", "Direct revision", "Stale rerun", "Not run"]


def test_streamlit_theme_config_is_light_and_secret_free():
    import tomllib
    config = tomllib.loads((ROOT / ".streamlit" / "config.toml").read_text())
    theme = config["theme"]
    assert theme["base"] == "light" and theme["primaryColor"] == "#3157D5"
    assert theme["backgroundColor"] == "#F7F8FA" and theme["textColor"] == "#111827"
    assert "key" not in (ROOT / ".streamlit" / "config.toml").read_text().lower().replace(
        "never put secrets here", "")
