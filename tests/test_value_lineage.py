"""Value lineage: RAW RESEARCHED VALUE -> optional USER OVERRIDE -> EFFECTIVE DEAL VALUE.

Regression case for the pattern found in the live Accenture -> EPAM demo: stale mapped
debt facts, a 2010 point-in-time share count, fresh cash, and documented overrides.
The facts run through the real SEC field builder. No network, no Gemini.
"""

import json
from pathlib import Path
import re
from dataclasses import replace

import pytest
from openpyxl import load_workbook

import ui_helpers as ui
from deal_analysis import (
    DEFAULT_SHARE_BASIS,
    SHARE_BASIS_DEFINITION_KEY,
    annotate_deal_context,
    build_synergy_context,
    build_transaction_structure_context,
    build_valuation_context,
    strip_input_lineage,
)
from agents import ICMemoWriter
from agents.registry import RISK, STRUCTURING
from memo import InvestmentCommitteeMemoGenerator
from models import load_deal
from orchestrator import APPROVED, DealTeamOrchestrator
from research import registry
from research.draft import create_draft
from research.errors import ResearchError
from research.models import (
    POINT_IN_TIME_SHARES,
    SHARE_BASES,
    WEIGHTED_BASIC_SHARES,
    WEIGHTED_DILUTED_SHARES,
    CompanyCandidate,
    ResearchStatus,
    ResearchValue,
)
from research.providers.sec.client import COMPANY_FACTS_URL, SUBMISSIONS_URL
from research.provenance import (
    SHARE_BASIS_LABEL,
    OverrideDetails,
    approve_draft,
    build_researched_deal,
    effective_label,
)
from research.quality import STALE_AFTER_DAYS, mark_if_stale, vintage_warning
from research_fixtures import FakeSECFetch, sec_provider
from excel_export import export_ma_workbook
from demo_state import SAMPLE_DEAL_PATH

US = registry.get_jurisdiction("US")
SAMPLE = load_deal(SAMPLE_DEAL_PATH)
ACN_CIK, EPAM_CIK = "0001467373", "0001352010"


def annual(concept_val_pairs, start, end, filed, taxonomy_unit="USD"):
    return {c: {"units": {taxonomy_unit: [{"start": start, "end": end, "val": v, "accn": f"a-{filed}",
                                           "fy": int(end[:4]), "fp": "FY", "form": "10-K",
                                           "filed": filed}]}}
            for c, v in concept_val_pairs}


def instant(end, val, filed, form="10-Q"):
    return {"end": end, "val": val, "accn": f"i-{filed}", "fy": int(filed[:4]), "fp": "Q3",
            "form": form, "filed": filed}


def acn_facts():
    us_gaap = {
        **annual([("Revenues", 69672977000), ("OperatingIncomeLoss", 10000000000),
                  ("Depreciation", 1000000000), ("AmortizationOfIntangibleAssets", 594049000),
                  ("NetIncomeLoss", 7678433000)], "2024-09-01", "2025-08-31", "2025-10-10"),
        "CashAndCashEquivalentsAtCarryingValue": {"units": {"USD": [
            instant("2026-05-31", 10165245000, "2026-06-18")]}},
        # the only mapped debt fact stops in 2024 (the live-demo pattern)
        "DebtLongtermAndShorttermCombinedAmount": {"units": {"USD": [
            instant("2024-08-31", 1024857000, "2024-10-10", form="10-K")]}},
        # a weighted-average EPS denominator: a different concept, never "shares outstanding"
        "WeightedAverageNumberOfSharesOutstandingBasic": {"units": {"shares": [
            {"start": "2026-03-01", "end": "2026-05-31", "val": 612213453, "accn": "w",
             "fy": 2026, "fp": "Q3", "form": "10-Q", "filed": "2026-06-18"}]}},
    }
    dei = {"EntityCommonStockSharesOutstanding": {"units": {"shares": [
        instant("2010-03-19", 637016143, "2010-03-26")]}}}
    return {"cik": 1467373, "entityName": "Accenture plc", "facts": {"us-gaap": us_gaap, "dei": dei}}


def epam_facts():
    us_gaap = {
        **annual([("RevenueFromContractWithCustomerExcludingAssessedTax", 5457056000),
                  ("OperatingIncomeLoss", 500000000), ("DepreciationAndAmortization", 144814000),
                  ("NetIncomeLoss", 377678000),
                  ("NetCashProvidedByUsedInOperatingActivities", 700000000),
                  ("PaymentsToAcquirePropertyPlantAndEquipment", 87309000)],
                 "2025-01-01", "2025-12-31", "2026-02-26"),
        "CashAndCashEquivalentsAtCarryingValue": {"units": {"USD": [
            instant("2026-06-30", 789397000, "2026-08-06")]}},
        "LongTermDebt": {"units": {"USD": [instant("2015-12-31", 35000000, "2016-02-20", form="10-K")]}},
    }
    dei = {"EntityCommonStockSharesOutstanding": {"units": {"shares": [
        instant("2026-07-31", 51586981, "2026-08-06")]}}}
    return {"cik": 1352010, "entityName": "EPAM Systems, Inc.", "facts": {"us-gaap": us_gaap, "dei": dei}}


def research(cik, name, facts):
    fetch = FakeSECFetch(overrides={
        SUBMISSIONS_URL.format(cik=cik): {"name": name, "fiscalYearEnd": "0831"},
        COMPANY_FACTS_URL.format(cik=cik): facts})
    candidate = CompanyCandidate(cik, name, None, "NYSE", "US", "CIK")
    return sec_provider(fetch).research_company(candidate)


PRICE = ResearchValue("share_price", ResearchStatus.MARKET, 174.47, "USD per share",
                      source_provider="Alpha Vantage (third-party market data)",
                      period_end="latest trading day 2026-09-28")

OVERRIDES = {"acquirer.debt": 5142.265, "acquirer.shares_outstanding": 612.213453,
             "target.debt": 25.0}
DETAILS = {
    "acquirer.debt": OverrideDetails(
        "Mapped XBRL debt fact is stale", "2026-05-31",
        "SEC 10-Q: LongTermDebtAndCapitalLeaseObligations 5,029.449 + DebtCurrent 112.816"),
    "acquirer.shares_outstanding": OverrideDetails(
        "No current point-in-time share count in CompanyFacts", "2026-05-31",
        "SEC 10-Q: WeightedAverageNumberOfSharesOutstandingBasic, quarter to 2026-05-31",
        WEIGHTED_BASIC_SHARES),
    "target.debt": OverrideDetails("Mapped XBRL debt fact is stale", "2026-06-30",
                                   "SEC 10-Q: LongTermDebtNoncurrent"),
}


def assumptions():
    from deal_input import form_from_deal
    from research.provenance import ASSUMPTION_PATHS
    values = {p: v for p, v in form_from_deal(SAMPLE).items() if p in ASSUMPTION_PATHS}
    values.update({"market_data.comparable_ev_ebitda_multiple": 8.2,
                   "market_data.precedent_ev_ebitda_multiple": 12.3,
                   "transaction_assumptions.offer_enterprise_value": 6300.0,
                   "transaction_assumptions.cash_percentage": 70.0,
                   "transaction_assumptions.stock_percentage": 30.0})
    return values


@pytest.fixture(scope="module")
def draft():
    return create_draft(US, research(ACN_CIK, "Accenture plc", acn_facts()),
                        research(EPAM_CIK, "EPAM Systems, Inc.", epam_facts()), PRICE)


@pytest.fixture(scope="module")
def provenance(draft):
    return approve_draft(draft, OVERRIDES, details=DETAILS)


@pytest.fixture(scope="module")
def deal(draft, provenance):
    return build_researched_deal(draft, provenance, assumptions(), "Hypothetical test deal")


@pytest.fixture(scope="module")
def contexts(deal, provenance):
    return DealTeamOrchestrator.build_deterministic_contexts(deal, provenance.lineage())


def approved_llm():
    class LLM:
        def generate(self, system_prompt, user_prompt):
            agent = re.match(r"You are the (.+?), ", system_prompt).group(1)
            if agent == "IC Partner":
                return json.dumps({"status": "approved", "summary": "Quality criteria met.",
                                   "revision_requests": [], "cross_workstream_issues": [],
                                   "unresolved_information_gaps": []})
            return f"{agent.upper()} OUTPUT"
    return LLM()


@pytest.fixture(scope="module")
def result(deal, provenance):
    return DealTeamOrchestrator(approved_llm()).run(deal, provenance.lineage())


# --- Stale point-in-time selection ------------------------------------------------------------

def test_stale_mapped_facts_are_flagged_not_used(draft):
    acn, epam = draft.acquirer.fields, draft.target.fields
    for value, end, number in ((acn["debt"], "2024-08-31", 1024.857),
                               (acn["shares_outstanding"], "2010-03-19", 637.016143),
                               (epam["debt"], "2015-12-31", 35.0)):
        assert value.status == ResearchStatus.STALE
        assert not value.available and value.observed           # kept, never used automatically
        assert (value.period_end, value.normalized_value) == (end, number)
        assert value.notes.startswith("Stale: as of")
    assert acn["cash"].status == ResearchStatus.REPORTED and acn["cash"].period_end == "2026-05-31"
    assert epam["cash"].status == ResearchStatus.REPORTED
    assert epam["shares_outstanding"].status == ResearchStatus.REPORTED  # fresh 2026 cover fact


def test_2010_share_fact_cannot_become_effective_in_2026(draft):
    with pytest.raises(ResearchError) as exc_info:
        approve_draft(draft)  # nothing entered for the stale rows
    message = str(exc_info.value)
    assert "Point-in-time shares outstanding is stale (as of 2010-03-19)" in message
    assert "Debt is stale (as of 2024-08-31)" in message and "Debt is stale (as of 2015-12-31)" in message


def test_weighted_average_shares_are_never_relabelled_as_shares_outstanding(draft):
    shares = draft.acquirer.fields["shares_outstanding"]
    assert shares.concept == "EntityCommonStockSharesOutstanding"
    assert "WeightedAverage" not in (shares.concept or "") and shares.normalized_value != 612.213453
    only_weighted = acn_facts()
    del only_weighted["facts"]["dei"]
    fields = research(ACN_CIK, "Accenture plc", only_weighted).fields
    assert fields["shares_outstanding"].status == ResearchStatus.MISSING


def test_latest_valid_point_in_time_context_is_selected_across_concepts():
    facts = acn_facts()
    facts["facts"]["us-gaap"]["CommonStockSharesOutstanding"] = {"units": {"shares": [
        instant("2026-05-31", 620000000, "2026-06-18")]}}
    shares = research(ACN_CIK, "Accenture plc", facts).fields["shares_outstanding"]
    # the 2010 cover-page fact no longer wins merely because its concept is listed first
    assert (shares.status, shares.concept, shares.period_end) == (
        ResearchStatus.REPORTED, "CommonStockSharesOutstanding", "2026-05-31")
    assert shares.normalized_value == 620.0


def test_staleness_threshold_is_about_twelve_months():
    value = ResearchValue("cash", ResearchStatus.REPORTED, 1.0, period_end="2025-01-01")
    assert mark_if_stale(value, "2026-01-02").status == ResearchStatus.REPORTED   # 366 days
    assert mark_if_stale(value, "2026-01-03").status == ResearchStatus.STALE      # 367 days
    assert STALE_AFTER_DAYS == 366
    assert mark_if_stale(value, None) is value


# --- Effective value resolution -----------------------------------------------------------------

def test_override_changes_effective_value_not_the_observation(draft, provenance):
    debt = provenance.approved["acquirer.debt"]
    assert (debt.value, debt.status, debt.as_of) == (5142.265, ResearchStatus.USER_OVERRIDE, "2026-05-31")
    assert debt.original is draft.acquirer.fields["debt"]                      # same object
    assert (debt.original.normalized_value, debt.original.period_end) == (1024.857, "2024-08-31")
    assert debt.override_source.startswith("SEC 10-Q")
    cash = provenance.approved["acquirer.cash"]
    assert (cash.status, cash.as_of, cash.value) == (ResearchStatus.REPORTED, "2026-05-31", 10165.245)


def test_share_proxy_override_is_labelled_as_a_proxy(provenance):
    shares = provenance.approved["acquirer.shares_outstanding"]
    assert shares.basis == WEIGHTED_BASIC_SHARES == provenance.share_basis
    assert "modelling proxy" in shares.basis and shares.basis != POINT_IN_TIME_SHARES


def test_invalid_override_date_or_basis_is_rejected(draft):
    with pytest.raises(ResearchError, match="YYYY-MM-DD"):
        approve_draft(draft, OVERRIDES, details={**DETAILS, "target.debt": OverrideDetails(as_of="30/06/2026")})
    with pytest.raises(ResearchError, match="unknown share basis"):
        approve_draft(draft, OVERRIDES, details={
            **DETAILS, "acquirer.shares_outstanding": OverrideDetails(basis="Shares outstanding")})


# --- One effective value everywhere --------------------------------------------------------------

def test_deal_holds_only_effective_values(deal):
    assert (deal.acquirer.debt, deal.acquirer.shares_outstanding, deal.target.debt) == (
        5142.265, 612.213453, 25.0)
    assert deal.acquirer.cash == 10165.245 and deal.target.cash == 789.397


def test_net_debt_uses_effective_debt_deterministically(deal, contexts):
    expected = 25.0 - 789.397
    assert contexts["valuation"]["target_net_debt"] == pytest.approx(expected)
    full = build_transaction_structure_context(deal, ndigits=None)
    assert full["calculation"]["target_net_debt"] == pytest.approx(expected)


def test_overview_uses_effective_values_with_vintage(deal, provenance, contexts):
    acquirer, target = ui.company_rows(deal, provenance, contexts)
    acquirer, target = dict(acquirer), dict(target)
    assert acquirer["Debt"] == "USD 5.1bn · as of 2026-05-31 · user override"
    assert acquirer["Cash"] == "USD 10.2bn · as of 2026-05-31"
    assert target["Debt"] == "USD 25.0m · as of 2026-06-30 · user override"
    assert "1.0bn" not in " ".join(acquirer.values())
    assert acquirer["Share basis"] == ("612.21m shares · " + WEIGHTED_BASIC_SHARES
                                       + " · as of 2026-05-31 · user override")


def test_research_view_shows_effective_value_and_preserves_the_original(draft, provenance):
    rows = {r[0]: r for r in ui.research_fact_rows(draft, "acquirer", provenance)}
    debt = rows["Debt"]
    assert (debt[1], debt[3], debt[4], debt[6]) == (
        "USD 5.1bn", "User override", "as of 2026-05-31", "Effective value")
    assert debt[7] == ("Original researched value (Debt): USD 1.0bn · Stale · SEC EDGAR · 10-K · "
                       "as of 2024-08-31")
    # a weighted-average proxy is presented as the acquirer share basis, never as
    # point-in-time shares outstanding; the raw SEC observation keeps its own concept label
    assert "Point-in-time shares outstanding" not in rows
    shares = rows["Acquirer share basis"]
    assert shares[1].startswith("612.213") and WEIGHTED_BASIC_SHARES in shares[5]
    assert shares[7].startswith(f"Definition: {WEIGHTED_BASIC_SHARES}. Original researched value "
                                "(Point-in-time shares outstanding): 637.016")
    assert "as of 2010-03-19" in shares[7]
    # before approval the stale observation is visibly marked as not used
    before = {r[0]: r for r in ui.research_fact_rows(draft, "acquirer")}
    assert (before["Debt"][3], before["Debt"][6]) == ("Stale", "Not used — stale observation")


def test_override_audit_lines_keep_both_values(provenance):
    lines = " ".join(ui.override_lines(provenance))
    assert "effective value USD 5.1bn (as of 2026-05-31)" in lines
    assert ("**Acquirer debt** — effective value USD 5.1bn" in lines
            and "original researched value (Debt): USD 1.0bn · Stale · SEC EDGAR · 10-K · "
                "as of 2024-08-31" in lines)
    assert "**Acquirer share basis** — effective value 612.213" in lines
    assert f"definition: {WEIGHTED_BASIC_SHARES}" in lines
    assert "original researched value (Point-in-time shares outstanding): 637.016" in lines


def test_agents_and_memo_receive_only_effective_values(deal, result):
    assert result.status == APPROVED
    context = InvestmentCommitteeMemoGenerator(approved_llm()).build_context(result, deal)
    assert context["deal"]["acquirer"]["debt"] == 5142.265
    assert context["deal"]["acquirer"]["shares_outstanding"] == 612.213453
    assert context["deal"]["target"]["debt"] == 25.0
    text = json.dumps(context)
    for raw in ("1024.857", "637.016143", "35.0,"):   # stale raw observations never enter
        assert raw not in text
    notes = context["deterministic_calculations"]["transaction_structure"]["calculation_notes"]
    assert notes["acquirer_share_basis_definition"] == WEIGHTED_BASIC_SHARES


def test_excel_uses_effective_values_and_keeps_lineage_in_comments(deal, result, provenance, tmp_path):
    path = export_ma_workbook(deal, result, None, tmp_path / "m.xlsx", deal_source="research",
                              input_sources=provenance.excel_sources())
    ws = load_workbook(path)["Assumptions"]
    rows = [(ws.cell(r, 1).value, ws.cell(r, 2)) for r in range(1, ws.max_row + 1)]
    cells = dict(rows)
    acquirer_debt, target_debt = [c for label, c in rows if label and label.startswith("Debt")]
    assert (acquirer_debt.value, target_debt.value) == (5142.265, 25)
    debt = acquirer_debt
    comment = debt.comment.text
    assert "EFFECTIVE VALUE 5142.265 — user override — as of 2026-05-31" in comment
    assert "Original researched value (kept for provenance, not used): 1024.857" in comment
    assert "as of 2024-08-31" in comment
    shares = cells["Acquirer share basis (millions shares)"]
    assert shares.value == 612.213453 and WEIGHTED_BASIC_SHARES in shares.comment.text
    assert "Shares outstanding (millions)" not in cells


# --- Seller ownership -------------------------------------------------------------------------------

def test_seller_ownership_reconciles_with_the_displayed_share_basis(deal, provenance, contexts):
    full = build_transaction_structure_context(deal, ndigits=None)
    base = full["transaction_terms"]["acquirer_shares_outstanding"]
    new = full["calculation"]["new_shares_issued"]
    assert base == deal.acquirer.shares_outstanding == 612.213453
    assert full["calculation"]["seller_ownership_of_pro_forma_shares"] == pytest.approx(
        new / (base + new), rel=1e-12)
    view = ui.share_basis_view(deal, contexts, provenance)
    assert view["value"] == ui.fmt_shares(base, deal) == "612.21m shares"
    assert view["definition"] == WEIGHTED_BASIC_SHARES
    tx = ui.transaction_view(deal, contexts)
    assert dict(tx["shares"])["Acquirer share basis used"] == "612.21m shares"
    assert tx["share_basis_definition"] == WEIGHTED_BASIC_SHARES


# --- Mixed vintages ---------------------------------------------------------------------------------

def test_no_vintage_warning_when_effective_dates_are_aligned(provenance):
    assert provenance.data_quality_warnings() == ()   # all effective dates within a quarter


def test_vintage_warning_when_dates_span_more_than_six_months(draft):
    details = {**DETAILS, "acquirer.debt": replace(DETAILS["acquirer.debt"], as_of="2025-08-31")}
    warnings = approve_draft(draft, OVERRIDES, details=details).data_quality_warnings()
    assert warnings == ("Accenture plc: balance-sheet inputs use materially different reporting "
                        "dates (Cash — 2026-05-31; Debt — 2025-08-31; Share basis — 2026-05-31). "
                        "Derived values that combine them therefore mix reporting vintages.",)


def test_target_vintage_warning_mentions_deterministic_net_debt(draft):
    details = {**DETAILS, "target.debt": replace(DETAILS["target.debt"], as_of="2025-06-30")}
    (warning,) = approve_draft(draft, OVERRIDES, details=details).data_quality_warnings()
    assert warning.startswith("EPAM Systems, Inc.: balance-sheet inputs use materially different")
    assert "Target net debt (debt − cash) is still computed deterministically" in warning


def test_vintage_threshold_is_about_six_months():
    assert vintage_warning("X", [("Cash", "2026-01-01"), ("Debt", "2026-07-02")]) is None  # 182 days
    assert vintage_warning("X", [("Cash", "2026-01-01"), ("Debt", "2026-07-04")])          # 184 days


def test_undated_override_is_flagged(draft):
    details = {**DETAILS, "target.debt": replace(DETAILS["target.debt"], as_of=None)}
    warnings = approve_draft(draft, OVERRIDES, details=details).data_quality_warnings()
    assert warnings == ("No as-of date was supplied for: EPAM Systems, Inc. debt. The reporting "
                        "vintage of these inputs cannot be checked.",)


def test_warning_flows_to_contexts_without_changing_arithmetic(draft, deal):
    details = {**DETAILS, "acquirer.debt": replace(DETAILS["acquirer.debt"], as_of="2024-08-31")}
    lineage = approve_draft(draft, OVERRIDES, details=details).lineage()
    annotated = DealTeamOrchestrator.build_deterministic_contexts(deal, lineage)
    plain = {"valuation": build_valuation_context(deal), "synergy": build_synergy_context(deal),
             "transaction_structure": build_transaction_structure_context(deal)}
    for name in plain:
        assert strip_input_lineage(annotated[name]) == strip_input_lineage(plain[name])
    for name in ("valuation", "transaction_structure"):
        assert annotated[name]["calculation_notes"]["data_quality_notes"] == list(lineage.data_quality_notes)
    assert ui.data_quality_notes(annotated) == lineage.data_quality_notes


# --- Regression: sample deal ------------------------------------------------------------------------

def test_sample_finance_is_unchanged_and_basis_is_labelled():
    contexts = DealTeamOrchestrator.build_deterministic_contexts(SAMPLE)
    plain = {"valuation": build_valuation_context(SAMPLE), "synergy": build_synergy_context(SAMPLE),
             "transaction_structure": build_transaction_structure_context(SAMPLE)}
    for name, context in plain.items():
        assert strip_input_lineage(contexts[name]) == context   # every number identical
    notes = contexts["transaction_structure"]["calculation_notes"]
    assert notes["acquirer_share_basis_definition"] == DEFAULT_SHARE_BASIS
    assert "data_quality_notes" not in notes
    assert ui.share_basis_view(SAMPLE, contexts)["definition"] == DEFAULT_SHARE_BASIS


# --- Streamlit: research review end to end -------------------------------------------------------

class FixtureProvider:
    """Stands in for the SEC provider, returning the regression research above."""

    results = {}

    def resolve_company(self, query):
        cik, name = {"ACN": (ACN_CIK, "Accenture plc"), "EPAM": (EPAM_CIK, "EPAM Systems, Inc.")}[query]
        return CompanyCandidate(cik, name, None, "NYSE", "US", "CIK")

    def research_company(self, candidate, force_refresh=False):
        facts = acn_facts() if candidate.provider_company_id == ACN_CIK else epam_facts()
        return research(candidate.provider_company_id, candidate.name, facts)


def test_app_research_review_shows_effective_and_original_values(monkeypatch):
    from streamlit.testing.v1 import AppTest

    import llm_client
    from config import load_settings

    monkeypatch.setattr(llm_client.GeminiClient, "__init__",
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError("Gemini")))
    monkeypatch.setattr("urllib.request.urlopen",
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError("network")))
    monkeypatch.setattr(ui, "load_settings", lambda: load_settings(env_file=None))
    monkeypatch.setitem(registry.PROVIDER_FACTORIES, "sec", FixtureProvider)
    monkeypatch.setitem(registry.MARKET_PRICE_FACTORIES, "US", lambda: None)  # price entered

    at = AppTest.from_file(str(Path(__file__).resolve().parent.parent / "app.py"),
                           default_timeout=120).run()
    at.sidebar.radio(key="deal_source").set_value(ui.RESEARCH_DEAL).run()
    at.text_input(key="research_q_acquirer").set_value("ACN")
    at.text_input(key="research_q_target").set_value("EPAM")
    next(b for b in at.button if b.label == "Research Companies").click().run()
    assert not at.exception
    page = " ".join(m.value for m in at.markdown)
    assert "Stale research" in page and "Not used — stale observation" in page
    # a stale observation is never prefilled as if it were current
    for path in OVERRIDES:
        assert at.number_input(key=f"rv_{path}").value is None

    for path, value in {**OVERRIDES, "acquirer.share_price": 174.47}.items():
        at.number_input(key=f"rv_{path}").set_value(value)
    for path, d in DETAILS.items():
        at.text_input(key=f"rr_{path}").set_value(d.reason)
        at.text_input(key=f"ro_{path}").set_value(d.as_of)
        at.text_input(key=f"rs_{path}").set_value(d.source)
    at.selectbox(key="rb_acquirer.shares_outstanding").set_value(WEIGHTED_BASIC_SHARES)
    at.number_input(key="ra_transaction_assumptions.offer_enterprise_value").set_value(6300.0)
    at.button(key="FormSubmitter:research_review_form-Build Deal").click().run()
    assert not at.exception and not at.error

    deal = at.session_state["researched_deal"]
    assert (deal.acquirer.debt, deal.acquirer.shares_outstanding, deal.target.debt) == (
        5142.265, 612.213453, 25.0)
    page = " ".join(m.value for m in at.markdown)
    assert "Effective value" in page
    assert "Original researched value (Debt): USD 1.0bn · Stale · SEC EDGAR · 10-K · as of 2024-08-31" in page
    assert "Acquirer share basis" in page and f"Definition: {WEIGHTED_BASIC_SHARES}." in page
    assert "USD 5.1bn · as of 2026-05-31 · user override" in page            # Overview
    assert "612.21m shares · " + WEIGHTED_BASIC_SHARES in page
    captions = " ".join(c.value for c in at.caption)
    assert ("Acquirer share basis used in the transaction model: 612.21m shares · "
            + WEIGHTED_BASIC_SHARES) in captions


# --- Share-basis terminology (labels and agent/memo context) ---------------------------------------

def approved_with_basis(draft, basis):
    details = {**DETAILS, "acquirer.shares_outstanding": replace(
        DETAILS["acquirer.shares_outstanding"], basis=basis)}
    return approve_draft(draft, OVERRIDES, details=details)


@pytest.mark.parametrize("basis, label", [
    (POINT_IN_TIME_SHARES, "Point-in-time shares outstanding"),
    (WEIGHTED_BASIC_SHARES, SHARE_BASIS_LABEL),
    (WEIGHTED_DILUTED_SHARES, SHARE_BASIS_LABEL),
])
def test_effective_share_label_follows_the_approved_basis(draft, basis, label):
    assert set(SHARE_BASES) == {POINT_IN_TIME_SHARES, WEIGHTED_BASIC_SHARES, WEIGHTED_DILUTED_SHARES}
    provenance = approved_with_basis(draft, basis)
    approved = provenance.approved["acquirer.shares_outstanding"]
    assert effective_label(approved, "Point-in-time shares outstanding") == label
    rows = [r for r in ui.research_fact_rows(draft, "acquirer", provenance) if r[1].startswith("612.213")]
    (row,) = rows
    assert row[0] == label
    assert row[7].startswith(f"Definition: {basis}. ")          # exact definition beneath the value
    # the raw SEC observation keeps its original concept label, whatever the approved basis
    assert "Original researched value (Point-in-time shares outstanding): 637.016" in row[7]
    assert draft.acquirer.fields["shares_outstanding"].concept == "EntityCommonStockSharesOutstanding"


def test_accepted_point_in_time_research_keeps_its_label():
    facts = acn_facts()
    facts["facts"]["us-gaap"]["CommonStockSharesOutstanding"] = {"units": {"shares": [
        instant("2026-05-31", 620000000, "2026-06-18")]}}
    fresh = create_draft(US, research(ACN_CIK, "Accenture plc", facts),
                         research(EPAM_CIK, "EPAM Systems, Inc.", epam_facts()), PRICE)
    overrides = {k: v for k, v in OVERRIDES.items() if k != "acquirer.shares_outstanding"}
    provenance = approve_draft(fresh, overrides, details=DETAILS)
    rows = {r[0]: r for r in ui.research_fact_rows(fresh, "acquirer", provenance)}
    row = rows["Point-in-time shares outstanding"]
    assert row[3] == "Reported" and row[7] == f"Definition: {POINT_IN_TIME_SHARES}."


def test_deterministic_contexts_and_all_agents_carry_the_definition(deal, provenance, result):
    contexts = result.deterministic_contexts
    notes = contexts["transaction_structure"]["calculation_notes"]
    assert notes["acquirer_share_basis_definition"] == WEIGHTED_BASIC_SHARES
    annotated = annotate_deal_context({"acquirer": {"shares_outstanding": 612.213453}},
                                      provenance.lineage())
    assert annotated["acquirer"]["shares_outstanding"] == 612.213453   # value untouched
    assert annotated["acquirer"][SHARE_BASIS_DEFINITION_KEY] == (
        "acquirer.shares_outstanding is the acquirer share basis used in the transaction "
        f"model. Definition: {WEIGHTED_BASIC_SHARES}.")


def test_agents_receive_the_definition_beside_the_share_count(deal, provenance):
    class Recorder:
        def __init__(self):
            self.inner, self.prompts = approved_llm(), []

        def generate(self, system_prompt, user_prompt):
            self.prompts.append((system_prompt, user_prompt))
            return self.inner.generate(system_prompt, user_prompt)

    llm = Recorder()
    DealTeamOrchestrator(llm).run(deal, provenance.lineage())
    expected = f"Definition: {WEIGHTED_BASIC_SHARES}."
    for agent in (STRUCTURING, RISK, "IC Partner"):
        prompts = [(sp, up) for sp, up in llm.prompts if sp.startswith(f"You are the {agent}, ")]
        assert prompts, agent
        for system_prompt, user_prompt in prompts:
            deal_data = user_prompt.split("=== DETERMINISTIC CALCULATIONS")[0]
            assert '"shares_outstanding": 612.213453' in deal_data and expected in deal_data
            assert "Never call it point-in-time shares outstanding" in system_prompt


def test_memo_writer_input_carries_the_definition(deal, result):
    context = InvestmentCommitteeMemoGenerator(approved_llm()).build_context(result, deal)
    acquirer = context["deal"]["acquirer"]
    assert acquirer["shares_outstanding"] == 612.213453
    assert acquirer[SHARE_BASIS_DEFINITION_KEY].endswith(f"Definition: {WEIGHTED_BASIC_SHARES}.")
    prompt = ICMemoWriter(approved_llm()).build_user_prompt(context)
    assert f"Definition: {WEIGHTED_BASIC_SHARES}." in prompt.split("=== DETERMINISTIC CALCULATIONS")[0]


def test_terminology_fix_changes_no_finance_number(deal, provenance, result):
    plain = {"valuation": build_valuation_context(deal), "synergy": build_synergy_context(deal),
             "transaction_structure": build_transaction_structure_context(deal)}
    for name, context in plain.items():
        assert strip_input_lineage(result.deterministic_contexts[name]) == context
    full = build_transaction_structure_context(deal, ndigits=None)
    base = full["transaction_terms"]["acquirer_shares_outstanding"]
    new = full["calculation"]["new_shares_issued"]
    assert full["calculation"]["seller_ownership_of_pro_forma_shares"] == pytest.approx(
        new / (base + new), rel=1e-12)
