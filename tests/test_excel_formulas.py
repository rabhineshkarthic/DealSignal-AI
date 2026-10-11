"""The restricted formula evaluator must fail loudly on anything it cannot verify."""

import pytest
from openpyxl import Workbook

from excel_formulas import EvaluationError, FormulaError, WorkbookEvaluator


def book(**cells):
    wb = Workbook()
    ws = wb.active
    ws.title = "Main"
    other = wb.create_sheet("Other Sheet")
    other["B1"] = 4
    for coordinate, value in cells.items():
        ws[coordinate] = value
    return wb


def evaluate(wb, coordinate="B9"):
    return WorkbookEvaluator(wb).value("Main", coordinate)


@pytest.mark.parametrize("formula, expected", [
    ("=1+2*3", 7.0), ("=(1+2)*3", 9.0), ("=2^3", 8.0), ("=-B1+10", 7.0),
    ("='Other Sheet'!$B$1*2", 8.0), ("=SUM(B1:B2)", 8), ("=MIN(B1:B2,1)", 1.0),
    ("=MAX(B1:B2)", 5), ("=AVERAGE(B1:B2)", 4.0), ("=ABS(-B1)", 3),
    ('=IF(B1>2,"big","small")', "big"), ('=IF(B1=3,"eq","ne")', "eq"),
    ("=1E-12*0", 0.0),
])
def test_supported_subset(formula, expected):
    assert evaluate(book(B1=3, B2=5, B9=formula)) == expected


def test_if_evaluates_only_the_chosen_branch():
    assert evaluate(book(B1=0, B9='=IF(B1=0,"n/a",SUM(1,1/B1))')) == "n/a"
    with pytest.raises(EvaluationError, match="DIV/0"):
        evaluate(book(B1=0, B9='=IF(B1=1,"n/a",1/B1)'))


@pytest.mark.parametrize("formula, message", [
    ("=VLOOKUP(B1,B1:B2,1)", "unsupported function"),
    ("='Missing'!B1", "missing sheet"),
    ("=B5", "empty cell"),
    ("=1/0", "DIV/0"),
    ("=1+", "expected"),
    ("=B1:B2", "range"),
    ("=#REF!+1", "#REF!"),
])
def test_failures_are_loud(formula, message):
    with pytest.raises(FormulaError, match=message):
        evaluate(book(B1=3, B2=5, B9=formula))


def test_circular_reference_detected():
    with pytest.raises(FormulaError, match="circular"):
        evaluate(book(B1="=B9", B9="=B1"))


def test_text_cells_are_not_evaluated_as_formulas():
    wb = book()
    ws = wb["Main"]
    ws["B9"] = "=1/0"
    ws["B9"].data_type = "s"
    assert WorkbookEvaluator(wb).evaluate_all() == {}
