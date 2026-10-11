"""Phase 7: custom deal input. No network, no Gemini, no API key.

Every test blocks real GeminiClient construction and outbound sockets, so deal
creation and deterministic analysis are proven to be fully offline.
"""

import io
import socket
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

import llm_client
import ui_helpers as ui
from config import GeminiSettings, load_settings
from deal_input import (
    FIELDS,
    DealInputError,
    build_deal,
    decimal_to_pct,
    form_from_deal,
    pct_to_decimal,
    validate_form,
)
from excel_export import SHEET_ORDER
from memo import InvestmentCommitteeMemoGenerator
from models import Deal, load_deal
from orchestrator import APPROVED, DealTeamOrchestrator
from test_memo import ScriptedLLM, approved, memo_json

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_PATH = ROOT / "data" / "sample_deal.json"


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    def no_client(*args, **kwargs):
        raise AssertionError("a real GeminiClient must not be created")

    def no_network(*args, **kwargs):
        raise AssertionError("no network access is allowed")

    monkeypatch.setattr(llm_client.GeminiClient, "__init__", no_client)
    monkeypatch.setattr(socket.socket, "connect", no_network)


@pytest.fixture
def sample():
    return load_deal(SAMPLE_PATH)


@pytest.fixture
def form(sample):
    return form_from_deal(sample)


def custom_form(form, **changes):
    values = dict(form)
    values.update({
        "deal_name": "Orion Systems acquiring Vega Data", "acquirer.name": "Orion Systems",
        "target.name": "Vega Data", "currency": "EUR", "target.ebitda": 22.5,
        "target.free_cash_flow": 12.0, "dcf_assumptions.wacc": 9.25,
        "transaction_assumptions.offer_enterprise_value": 260.0,
        "transaction_assumptions.cash_percentage": 55.0,
        "transaction_assumptions.stock_percentage": 45.0,
    })
    values.update(changes)
    return values


# --- Percent conversion ----------------------------------------------------------------------

@pytest.mark.parametrize("percent, decimal", [
    (10.0, 0.1), (5.5, 0.055), (2.5, 0.025), (12.3, 0.123), (33.3, 0.333), (0.7, 0.007),
    (100, 1.0), (0, 0.0), (-3.5, -0.035),
])
def test_percent_to_decimal_is_exact(percent, decimal):
    assert pct_to_decimal(percent) == decimal
    assert repr(pct_to_decimal(percent)) == repr(decimal)  # no float residue
    assert decimal_to_pct(decimal) == float(percent)


def test_naive_division_would_have_left_residue():
    assert repr(12.3 / 100) != "0.123"  # why Decimal conversion is needed
    assert repr(pct_to_decimal(12.3)) == "0.123"


# --- Building the existing Deal type ------------------------------------------------------------

def test_sample_round_trips_exactly(sample, form):
    assert form["dcf_assumptions.wacc"] == 10.0
    assert form["transaction_assumptions.new_debt_interest_rate"] == 5.5
    assert build_deal(form) == sample


def test_custom_deal_is_the_existing_deal_type(form):
    deal = build_deal(custom_form(form))
    assert type(deal) is Deal
    assert deal.acquirer.name == "Orion Systems" and deal.target.name == "Vega Data"
    assert deal.dcf_assumptions.wacc == 0.0925
    assert deal.transaction_assumptions.cash_percentage == 0.55
    assert isinstance(deal.dcf_assumptions.forecast_years, int)


def test_names_are_trimmed(form):
    deal = build_deal(custom_form(form, **{"target.name": "  Vega Data  "}))
    assert deal.target.name == "Vega Data"


def test_deterministic_contexts_work_for_custom_deal(form):
    deal = build_deal(custom_form(form))
    contexts = DealTeamOrchestrator.build_deterministic_contexts(deal)
    assert contexts["valuation"]["comparable_company"]["enterprise_value"] == pytest.approx(22.5 * 10.5)
    assert contexts["transaction_structure"]["calculation"]["offer_equity_value"] == pytest.approx(240.0)
    preview = dict(ui.deal_preview(deal, contexts))
    assert preview["Offer EV"] == "EUR 260.0m"
    assert preview["Offer equity value"] == "EUR 240.0m"
    assert set(preview) == {"Offer EV", "Offer equity value", "Comparable companies EV",
                            "Precedent transactions EV", "DCF EV", "Valuation range",
                            "Net synergy value", "Cash / stock consideration"}


# --- Validation ------------------------------------------------------------------------------------

def errors_for(form, **changes):
    with pytest.raises(DealInputError) as exc_info:
        build_deal(custom_form(form, **changes))
    return exc_info.value.errors


@pytest.mark.parametrize("cash, stock", [(60.0, 50.0), (60.0, 30.0), (100.0, 0.1)])
def test_cash_plus_stock_must_equal_100(form, cash, stock):
    errors = errors_for(form, **{"transaction_assumptions.cash_percentage": cash,
                                 "transaction_assumptions.stock_percentage": stock})
    assert any("Cash % + Stock % must equal 100%" in e for e in errors)


def test_all_cash_and_all_stock_mixes_allowed(form):
    for cash in (100.0, 0.0):
        build_deal(custom_form(form, **{"transaction_assumptions.cash_percentage": cash,
                                        "transaction_assumptions.stock_percentage": 100.0 - cash}))


@pytest.mark.parametrize("terminal", [9.25, 12.0])
def test_wacc_must_exceed_terminal_growth(form, terminal):
    errors = errors_for(form, **{"dcf_assumptions.terminal_growth_rate": terminal})
    assert "WACC must be greater than the terminal growth rate." in errors


@pytest.mark.parametrize("field", ["acquirer.name", "target.name", "deal_name", "currency", "units"])
@pytest.mark.parametrize("bad", ["", "   ", None])
def test_text_fields_required(form, field, bad):
    errors = errors_for(form, **{field: bad})
    assert any("must not be empty" in e for e in errors)


@pytest.mark.parametrize("field, value, message", [
    ("acquirer.shares_outstanding", 0.0, "greater than zero"),
    ("acquirer.share_price", -1.0, "greater than zero"),
    ("market_data.comparable_ev_ebitda_multiple", 0.0, "greater than zero"),
    ("target.debt", -5.0, "must not be negative"),
    ("transaction_assumptions.offer_enterprise_value", -1.0, "must not be negative"),
    ("synergy_assumptions.integration_cost", -1.0, "must not be negative"),
    ("dcf_assumptions.forecast_years", 0, "whole number"),
    ("synergy_assumptions.realization_years", 2.5, "whole number"),
    ("dcf_assumptions.wacc", 0.0, "must be in"),
    ("synergy_assumptions.tax_rate", 100.0, "must be in"),
    ("synergy_assumptions.discount_rate", 150.0, "must be in"),
    ("transaction_assumptions.new_debt_interest_rate", -1.0, "must be in"),
    ("dcf_assumptions.fcf_growth_rate", -100.0, "greater than -100%"),
    ("target.revenue", float("nan"), "finite number"),
    ("acquirer.cash", "lots", "finite number"),
])
def test_field_validation(form, field, value, message):
    errors = errors_for(form, **{field: value})
    assert any(message in e for e in errors), errors


def test_all_errors_reported_together(form):
    errors = errors_for(form, **{"target.name": "", "dcf_assumptions.terminal_growth_rate": 20.0,
                                 "transaction_assumptions.cash_percentage": 90.0})
    assert len(errors) == 3


def test_model_validation_still_applies(form):
    """Rules enforced by valuation.py surface as input errors (not duplicated here)."""
    errors = errors_for(form, **{"transaction_assumptions.offer_enterprise_value": 10.0})
    assert errors == ["The model rejected these inputs: Net debt (20.0) exceeds enterprise "
                      "value (10.0); equity value would be negative"]


def test_missing_fields_reported(form):
    values = dict(form)
    del values["target.ebitda"]
    assert validate_form(values)[0].startswith("Missing fields: EBITDA")


def test_every_field_has_a_form_entry(form):
    assert set(form) == set(FIELDS)


# --- Stale-analysis protection --------------------------------------------------------------------

def test_switching_deal_clears_stale_analysis(sample, form):
    state = {"active_deal": sample, "deal_team_result": "R", "partial_result": "P",
             "memo_result": "M", "last_error": "E", "excel_bytes": b"X"}
    assert ui.reset_if_deal_changed(state, sample) is False
    assert state["deal_team_result"] == "R"  # same deal: nothing cleared
    custom = build_deal(custom_form(form))
    assert ui.reset_if_deal_changed(state, custom) is True
    assert all(state[k] is None for k in ui.ANALYSIS_KEYS)
    assert state["active_deal"] == custom


def test_editing_a_deal_clears_stale_analysis(form):
    first = build_deal(custom_form(form))
    edited = build_deal(custom_form(form, **{"target.ebitda": 30.0}))
    state = {"active_deal": first, "deal_team_result": "R", "memo_result": "M"}
    assert ui.reset_if_deal_changed(state, edited) is True
    assert state["deal_team_result"] is None and state["memo_result"] is None


def test_selecting_custom_source_without_deal_clears_analysis(sample):
    state = {"active_deal": sample, "deal_team_result": "R"}
    assert ui.reset_if_deal_changed(state, None) is True
    assert state["deal_team_result"] is None


# --- Export and live pipeline (mocked) --------------------------------------------------------------

def test_custom_deal_exports_to_excel(form):
    deal = build_deal(custom_form(form, **{"acquirer.name": "Orion/Systems Inc.",
                                           "target.name": "Vega: Data"}))
    llm = ScriptedLLM([approved()], memo_response=memo_json())
    result = DealTeamOrchestrator(llm).run(deal)
    memo = InvestmentCommitteeMemoGenerator(llm).generate(result, deal)
    data = ui.workbook_bytes(deal, result, memo, ui.CUSTOM_SOURCE_LABEL)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == list(SHEET_ORDER)
    revenue = next(c for row in wb["Assumptions"].iter_rows() for c in row
                   if c.comment and "acquirer.revenue" in c.comment.text)
    assert revenue.comment.text == "Source: custom deal input (Streamlit form) — acquirer.revenue"
    assert ui.workbook_filename(deal) == "OrionSystems_Vega_MA_Model.xlsx"


def test_filenames_are_safe_and_sample_unchanged(sample):
    assert ui.workbook_filename(sample) == "Northstar_Atlas_MA_Model.xlsx"
    odd = replace(sample, acquirer=replace(sample.acquirer, name="../../etc"),
                  target=replace(sample.target, name="***"))
    name = ui.workbook_filename(odd)
    assert "/" not in name and ".." not in name and name == "etc_Deal_MA_Model.xlsx"


def test_custom_deal_runs_through_live_pipeline_with_mocked_llm(form):
    deal = build_deal(custom_form(form))
    llm = ScriptedLLM([approved()], memo_response=memo_json())
    result = DealTeamOrchestrator(llm).run(deal)
    assert result.status == APPROVED
    assert result.deterministic_contexts == DealTeamOrchestrator.build_deterministic_contexts(deal)
    memo = InvestmentCommitteeMemoGenerator(llm).generate(result, deal)
    assert memo.deal_name == "Orion Systems acquiring Vega Data"


def test_sample_deal_file_unchanged(sample):
    assert sample.deal_name == "Northstar Software acquiring Atlas Analytics"
    assert sample.transaction_assumptions.offer_enterprise_value == 205.0
    assert sample.dcf_assumptions.wacc == 0.1


# --- Streamlit AppTest ----------------------------------------------------------------------------------

def app_test():
    from streamlit.testing.v1 import AppTest
    return AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)


@pytest.fixture
def unconfigured(monkeypatch):
    monkeypatch.setattr(ui, "load_settings", lambda: load_settings(env_file=None))


def submit(at):
    at.button(key="FormSubmitter:custom_deal_form-Create Deal").click().run()


def test_app_sample_deal_is_default(unconfigured):
    at = app_test().run()
    assert not at.exception
    assert at.session_state["deal_source"] == ui.SAMPLE_DEAL
    assert at.session_state["mode"] == ui.OFFLINE_DEMO


def test_app_create_custom_deal_offline(unconfigured):
    at = app_test().run()
    at.session_state["deal_team_result"] = "STALE"
    at.sidebar.radio(key="deal_source").set_value(ui.CUSTOM_DEAL).run()
    assert not at.exception
    assert at.session_state["deal_team_result"] is None  # stale analysis cleared
    assert any(ui.PUBLIC_RESEARCH_HINT in m.value for m in at.markdown)
    assert any(b.label == "Go to Public Company Research" and not b.disabled for b in at.button)
    at.text_input(key="form_acquirer.name").set_value("Orion Systems")
    at.text_input(key="form_target.name").set_value("Vega Data")
    at.number_input(key="form_target.ebitda").set_value(22.5)
    at.number_input(key="form_dcf_assumptions.wacc").set_value(9.25)
    submit(at)
    assert not at.exception
    deal = at.session_state["custom_deal"]
    assert deal.target.ebitda == 22.5 and deal.dcf_assumptions.wacc == 0.0925
    text = " ".join(m.value for m in at.markdown)
    assert "Deal created · Orion Systems → Vega Data" in text
    assert "USD 236.2m" in text or "USD 236.3m" in text  # 22.5 x 10.5 comps EV
    assert "User-defined transaction" in text


def test_app_invalid_custom_deal_shows_errors(unconfigured):
    at = app_test().run()
    at.sidebar.radio(key="deal_source").set_value(ui.CUSTOM_DEAL).run()
    at.number_input(key="form_transaction_assumptions.cash_percentage").set_value(70.0)
    submit(at)
    assert not at.exception
    assert at.session_state["custom_deal"] is None
    assert any("Cash % + Stock % must equal 100%" in e.value for e in at.error)


def test_app_live_analysis_on_custom_deal(monkeypatch):
    monkeypatch.setattr(ui, "live_settings_summary", lambda: ui.LiveSettingsSummary(
        True, "primary-model", (), None, GeminiSettings(api_key="k", model="primary-model")))
    monkeypatch.setattr(ui, "create_live_client",
                        lambda settings: ScriptedLLM([approved()], memo_response=memo_json()))
    at = app_test().run()
    at.sidebar.radio(key="mode").set_value(ui.LIVE_ANALYSIS).run()
    at.sidebar.radio(key="deal_source").set_value(ui.CUSTOM_DEAL).run()
    assert next(b for b in at.sidebar.button if b.label == "Run Full Deal Team").disabled
    at.text_input(key="form_target.name").set_value("Vega Data")
    at.number_input(key="form_target.ebitda").set_value(22.5)
    submit(at)
    at.button(key="run_from_preview").click().run()
    assert not at.exception
    result = at.session_state["deal_team_result"]
    deal = at.session_state["custom_deal"]
    assert result.status == APPROVED
    assert result.deterministic_contexts == DealTeamOrchestrator.build_deterministic_contexts(deal)
    # switching back to the sample deal clears the custom deal's analysis
    at.sidebar.radio(key="deal_source").set_value(ui.SAMPLE_DEAL).run()
    assert at.session_state["deal_team_result"] is None


def test_app_go_to_public_company_research_switches_source(unconfigured):
    at = app_test().run()
    at.sidebar.radio(key="deal_source").set_value(ui.CUSTOM_DEAL).run()
    next(b for b in at.button if b.label == "Go to Public Company Research").click().run()
    assert not at.exception
    assert at.session_state["deal_source"] == ui.RESEARCH_DEAL
    assert [s.label for s in at.selectbox] == ["Research jurisdiction"]  # research flow shown
    assert not any(b.label == "Go to Public Company Research" for b in at.button)


def test_deal_sources_are_described_and_no_stale_research_copy():
    assert set(ui.DEAL_SOURCE_DESCRIPTIONS) == set(ui.DEAL_SOURCES)
    stale = ["coming next", "not available yet", "research is not implemented",
             "Real Company Research"]
    for rel in ["app.py", "ui_helpers.py", "ui_components.py", "research_ui.py", "README.md"]:
        text = (ROOT / rel).read_text()
        assert not [p for p in stale if p.lower() in text.lower()], rel
