"""Phase 8.2: loss-making targets and EV/Revenue multiples. No network, no Gemini."""

import io
import json
import socket
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

import llm_client
import ui_helpers as ui
from agents import ValuationAnalyst
from deal_analysis import (
    build_synergy_context,
    build_transaction_structure_context,
    build_valuation_context,
)
from deal_input import DealInputError, build_deal, form_from_deal
from excel_formulas import WorkbookEvaluator
from memo import InvestmentCommitteeMemoGenerator
from models import Deal, deal_to_context, load_deal
from orchestrator import APPROVED, DealTeamOrchestrator
from research import approve_draft, build_researched_deal, registry
from research.draft import create_draft
from research.market_data import missing_price
from research.models import CompanyCandidate, ResearchStatus
from research.provenance import ASSUMPTION_PATHS
from research.providers.sec.client import COMPANY_FACTS_URL, SUBMISSIONS_URL
from research_fixtures import FakeSECFetch, sec_provider
from test_memo import ScriptedLLM, approved, memo_json
from valuation import (
    NON_POSITIVE_EBITDA_REASON,
    TOO_FEW_METHODS_NOTE,
    comparable_company_revenue_valuation,
    comparable_ebitda_method,
    precedent_ebitda_method,
    precedent_transaction_revenue_valuation,
    select_valuation_methods,
    summarize_valuation,
)

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = load_deal(ROOT / "data" / "sample_deal.json")
SNAPSHOT = json.loads((ROOT / "tests" / "fixtures" / "sample_contexts_before_phase_8_2.json").read_text())
REVENUE_MULTIPLES = {"market_data.comparable_ev_revenue_multiple": 8.0,
                     "market_data.precedent_ev_revenue_multiple": 10.0}
NO_EBITDA_MULTIPLES = {"market_data.comparable_ev_ebitda_multiple": None,
                       "market_data.precedent_ev_ebitda_multiple": None}


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def blocked(*args, **kwargs):
        raise AssertionError("network or Gemini access attempted")

    monkeypatch.setattr(socket.socket, "connect", blocked)
    monkeypatch.setattr("urllib.request.urlopen", blocked)
    monkeypatch.setattr(llm_client.GeminiClient, "__init__", blocked)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)


def loss_deal(ebitda=-45.0, **changes) -> Deal:
    values = {**form_from_deal(SAMPLE), **NO_EBITDA_MULTIPLES, **REVENUE_MULTIPLES,
              "target.ebitda": ebitda, **changes}
    return build_deal(values)


def leaves(obj, path=()):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield from leaves(v, path + (k,))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from leaves(v, path + (i,))
    else:
        yield path, obj


def get(obj, path):
    for key in path:
        obj = obj[key]
    return obj


# --- Backward compatibility ------------------------------------------------------------------

def test_sample_outputs_identical_to_pre_phase_snapshot():
    now = {
        "full_precision": {"valuation": build_valuation_context(SAMPLE, ndigits=None),
                           "synergy": build_synergy_context(SAMPLE, ndigits=None),
                           "transaction_structure": build_transaction_structure_context(SAMPLE, ndigits=None)},
        "display": {"valuation": build_valuation_context(SAMPLE),
                    "synergy": build_synergy_context(SAMPLE),
                    "transaction_structure": build_transaction_structure_context(SAMPLE)},
    }
    checked = 0
    for path, value in leaves(SNAPSHOT):
        assert get(now, path) == value, path  # exact equality, every value
        checked += 1
    assert checked == 328


def test_sample_still_uses_ev_ebitda_and_unchanged_deal_data():
    ctx = build_valuation_context(SAMPLE)
    assert ctx["valuation_basis"]["primary_multiple_basis"] == "EV/EBITDA"
    assert [m["method"] for m in ctx["valuation_summary"]["methods"]] == [
        "Comparable Companies", "Precedent Transactions", "Discounted Cash Flow"]
    assert "comparable_company_revenue" not in ctx
    raw = json.loads((ROOT / "data" / "sample_deal.json").read_text())
    raw.pop("disclaimer")
    assert deal_to_context(SAMPLE) == raw


# --- Accepting non-positive EBITDA -----------------------------------------------------------------

@pytest.mark.parametrize("ebitda", [0.0, -45.0, -0.5])
def test_non_positive_target_ebitda_accepted_and_preserved(ebitda):
    deal = loss_deal(ebitda)
    assert deal.target.ebitda == ebitda  # never forced positive, zero-ed, or MISSING


def test_negative_acquirer_ebitda_accepted():
    assert loss_deal(**{"acquirer.ebitda": -10.0}).acquirer.ebitda == -10.0


def test_revenue_multiples_required_for_non_positive_ebitda():
    values = {**form_from_deal(SAMPLE), **NO_EBITDA_MULTIPLES, "target.ebitda": -45.0}
    with pytest.raises(DealInputError) as exc_info:
        build_deal(values)
    assert len(exc_info.value.errors) == 2
    assert all("is required: target EBITDA is non-positive" in e for e in exc_info.value.errors)


def test_ebitda_multiples_still_required_for_positive_ebitda():
    with pytest.raises(DealInputError, match="required because target EBITDA is positive"):
        build_deal({**form_from_deal(SAMPLE), **NO_EBITDA_MULTIPLES})


@pytest.mark.parametrize("changes, message", [
    ({"market_data.comparable_ev_revenue_multiple": 0.0}, "greater than zero"),
    ({"target.revenue": 0.0}, "greater than zero for EV/Revenue"),
])
def test_revenue_valuation_inputs_validated(changes, message):
    with pytest.raises(DealInputError, match=message):
        loss_deal(**changes)


# --- EV/EBITDA applicability ------------------------------------------------------------------------

@pytest.mark.parametrize("ebitda", [0.0, -45.0])
def test_ev_ebitda_not_applicable_and_no_negative_ev(ebitda):
    for method in (comparable_ebitda_method(ebitda, 10.5), precedent_ebitda_method(ebitda, 12.0)):
        assert method["applicable"] is False and method["enterprise_value"] is None
        assert method["reason"] == NON_POSITIVE_EBITDA_REASON


def test_ev_ebitda_formula_unchanged_when_applicable():
    method = comparable_ebitda_method(18.0, 10.5)
    assert method["applicable"] and method["enterprise_value"] == 189.0


# --- EV/Revenue methods --------------------------------------------------------------------------------

def test_revenue_comps_and_precedents():
    comps = comparable_company_revenue_valuation(120.0, 8.0)
    precs = precedent_transaction_revenue_valuation(120.0, 10.0)
    assert (comps["enterprise_value"], precs["enterprise_value"]) == (960.0, 1200.0)
    assert comps["basis"] == precs["basis"] == "EV/Revenue"
    assert comps["method"] == "Comparable Companies (EV/Revenue)"


@pytest.mark.parametrize("revenue, multiple", [(0.0, 8.0), (120.0, 0.0), (-1.0, 8.0)])
def test_revenue_method_rejects_non_positive_inputs(revenue, multiple):
    with pytest.raises(ValueError):
        comparable_company_revenue_valuation(revenue, multiple)


# --- Range from applicable methods only ----------------------------------------------------------------

def test_range_excludes_not_applicable_methods():
    ctx = build_valuation_context(loss_deal(), ndigits=None)
    summary = ctx["valuation_summary"]
    assert set(summary["by_method"]) == {"Comparable Companies (EV/Revenue)",
                                         "Precedent Transactions (EV/Revenue)",
                                         "Discounted Cash Flow"}
    assert summary["low"] == pytest.approx(172.024, abs=1e-3) and summary["high"] == 1200.0
    assert summary["range_available"] and summary["low_method"] == "Discounted Cash Flow"
    assert set(ctx["transaction"]["offer_vs_method"]) == set(summary["by_method"])
    assert [m["applicable"] for m in summary["methods"]] == [False, False, True, True, True]


def test_too_few_methods_does_not_fabricate_a_range():
    methods = select_valuation_methods(-45.0, [comparable_ebitda_method(-45.0, 10.5)], [],
                                       {"enterprise_value": 172.0})
    summary = summarize_valuation(methods)
    assert summary["range_available"] is False and summary["range_note"] == TOO_FEW_METHODS_NOTE
    assert summary["low"] is None and summary["high"] is None and summary["average"] is None


def test_revenue_multiples_for_profitable_target_are_supplementary():
    deal = replace(SAMPLE, market_data=replace(SAMPLE.market_data,
                                               comparable_ev_revenue_multiple=8.0))
    ctx = build_valuation_context(deal, ndigits=None)
    base = build_valuation_context(SAMPLE, ndigits=None)
    assert ctx["valuation_summary"]["by_method"] == base["valuation_summary"]["by_method"]
    assert ctx["valuation_summary"]["low"] == base["valuation_summary"]["low"]
    extra = ctx["valuation_summary"]["methods"][2]
    assert extra["method"] == "Comparable Companies (EV/Revenue)"
    assert extra["applicable"] and not extra["included_in_range"]


# --- DCF with negative EBITDA ------------------------------------------------------------------------------

def test_dcf_runs_with_negative_ebitda():
    ctx = build_valuation_context(loss_deal(), ndigits=None)
    assert ctx["dcf"]["enterprise_value"] == pytest.approx(172.024, abs=1e-3)


def test_non_positive_fcf_surfaced_separately_from_ebitda():
    with pytest.raises(DealInputError, match="Free cash flow must be greater than zero"):
        loss_deal(**{"target.free_cash_flow": -5.0})


# --- Agent context ---------------------------------------------------------------------------------------------

def test_valuation_analyst_context_labels_methods_not_applicable():
    deal = loss_deal()
    ctx = build_valuation_context(deal)
    captured = {}

    class LLM:
        def generate(self, system_prompt, user_prompt):
            captured.update(system=system_prompt, user=user_prompt)
            return "Valuation reviewed."

    ValuationAnalyst(LLM()).run({"deal": deal_to_context(deal), "deterministic_calculations": ctx})
    assert '"primary_multiple_basis": "EV/Revenue"' in captured["user"]
    assert '"target_ebitda_status": "negative"' in captured["user"]
    assert NON_POSITIVE_EBITDA_REASON in captured["user"]
    assert '"applicable": false' in captured["user"]
    assert "not missing evidence or a calculation failure" in captured["user"]
    assert "Never describe\n  it as missing evidence, a data gap, or a calculation failure" in captured["system"]


# --- Snowflake-style researched target ----------------------------------------------------------------------------

def sigma_research():
    facts = json.loads((ROOT / "tests" / "fixtures" / "sec" / "companyfacts_0000000505.json").read_text())
    fetch = FakeSECFetch(overrides={
        SUBMISSIONS_URL.format(cik="0000000505"): {"name": "SIGMA CLOUD DATA INC"},
        COMPANY_FACTS_URL.format(cik="0000000505"): facts})
    provider = sec_provider(fetch)
    acquirer = provider.research_company(provider.resolve_company("ALPH"))
    target = provider.research_company(CompanyCandidate("0000000505", "SIGMA CLOUD DATA INC",
                                                        "SIGM", "NYSE", "US", "CIK"))
    return acquirer, target


def test_snowflake_style_target_researched_with_negative_ebitda():
    _, target = sigma_research()
    ebitda = target.fields["ebitda"]
    assert ebitda.status == ResearchStatus.DERIVED and ebitda.normalized_value == -1010.0
    assert target.fields["free_cash_flow"].normalized_value == 910.0
    assert target.fields["debt"].normalized_value == 2268.0


def build_sigma_deal():
    acquirer, target = sigma_research()
    draft = create_draft(registry.get_jurisdiction("US"), acquirer, target,
                         missing_price("No market price in tests."))
    provenance = approve_draft(draft, {"acquirer.share_price": 412.37})
    assumptions = {p: v for p, v in form_from_deal(SAMPLE).items() if p in ASSUMPTION_PATHS}
    assumptions.update({"market_data.comparable_ev_ebitda_multiple": None,
                        "market_data.precedent_ev_ebitda_multiple": None,
                        "market_data.comparable_ev_revenue_multiple": 12.0,
                        "market_data.precedent_ev_revenue_multiple": 14.0,
                        "transaction_assumptions.offer_enterprise_value": 60000.0})
    return build_researched_deal(draft, provenance,
                                 {k: v for k, v in assumptions.items() if v is not None})


def test_snowflake_style_deal_reaches_deterministic_preview():
    deal = build_sigma_deal()
    assert deal.target.ebitda == -1010.0  # kept exactly
    contexts = DealTeamOrchestrator.build_deterministic_contexts(deal)
    preview = dict(ui.deal_preview(deal, contexts))
    assert preview["Comparable companies EV"] == "Not applicable — target EBITDA non-positive"
    assert preview["Precedent transactions EV"] == "Not applicable — target EBITDA non-positive"
    assert preview["Comparable companies EV (EV/Revenue)"] == "USD 55.2bn"   # 4,600 x 12 = 55,200m
    assert preview["Precedent transactions EV (EV/Revenue)"] == "USD 64.4bn"  # 4,600 x 14 = 64,400m
    valuation = contexts["valuation"]
    assert valuation["comparable_company_revenue"]["enterprise_value"] == 55200.0  # still millions
    assert valuation["precedent_transaction_revenue"]["enterprise_value"] == 64400.0
    assert preview["Valuation range"].startswith("USD ")
    assert contexts["valuation"]["valuation_basis"]["primary_multiple_basis"] == "EV/Revenue"


# --- Excel -------------------------------------------------------------------------------------------------------------

def test_excel_shows_not_applicable_and_revenue_methods():
    deal = loss_deal()
    llm = ScriptedLLM([approved()], memo_response=memo_json())
    result = DealTeamOrchestrator(llm).run(deal)
    memo = InvestmentCommitteeMemoGenerator(llm).generate(result, deal)
    wb = load_workbook(io.BytesIO(ui.workbook_bytes(deal, result, memo)))
    ws = wb["Valuation"]
    rows = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(4, ws.max_row + 1)
            if ws.cell(r, 1).value}
    assert rows["Comparable Companies"] == "NOT APPLICABLE"
    assert rows["Precedent Transactions"] == "NOT APPLICABLE"
    ev = WorkbookEvaluator(wb)
    ev.evaluate_all()  # every formula still evaluates
    values = {ws.cell(r, 1).value: ev.value("Valuation", f"B{r}") for r in range(4, 18)
              if ws.cell(r, 1).value and ws.cell(r, 2).data_type == "f"}
    assert values["Comparable Companies (EV/Revenue)"] == pytest.approx(960.0)
    assert values["Precedent Transactions (EV/Revenue)"] == pytest.approx(1200.0)
    assert values["Low"] == pytest.approx(172.024, abs=1e-3) and values["High"] == pytest.approx(1200.0)
    assert 0 not in [ws.cell(r, 2).value for r in range(4, ws.max_row + 1)]  # no zero written
    summary = wb["IC Summary"]
    snapshot = {summary.cell(r, 1).value: summary.cell(r, 2).value for r in range(4, summary.max_row + 1)}
    assert snapshot["Comparable Companies EV"] == "NOT APPLICABLE"


# --- Streamlit ------------------------------------------------------------------------------------------------------

def test_app_valuation_tab_notice_for_negative_ebitda(monkeypatch):
    from config import load_settings
    monkeypatch.setattr(ui, "load_settings", lambda: load_settings(env_file=None))
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120).run()
    at.sidebar.radio(key="deal_source").set_value(ui.CUSTOM_DEAL).run()
    at.number_input(key="form_target.ebitda").set_value(-45.0)
    at.number_input(key="form_market_data.comparable_ev_revenue_multiple").set_value(8.0)
    at.number_input(key="form_market_data.precedent_ev_revenue_multiple").set_value(10.0)
    at.button(key="FormSubmitter:custom_deal_form-Create Deal").click().run()
    assert not at.exception
    assert at.session_state["custom_deal"].target.ebitda == -45.0
    page = " ".join(m.value for m in at.markdown)
    assert ui.NON_POSITIVE_EBITDA_NOTICE in page
    assert "Not applicable" in page and NON_POSITIVE_EBITDA_REASON in page
    assert "(USD 45.0m)" in page  # negative EBITDA shown normally, not as an error
    assert not at.error
