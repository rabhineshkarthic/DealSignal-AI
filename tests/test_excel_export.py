"""Phase 5 Excel export tests. Real .xlsx round-trips; scripted fake LLM only.

No Gemini, no API key, no network, no Excel application.
"""

import ast
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

import excel_export
from excel_export import (
    FINANCING_WARNING,
    NO_MEMO_NOTE,
    SHEET_ORDER,
    SYNERGY_CASE_CAUTION,
    ExcelExportError,
    build_workbook,
    export_ma_workbook,
)
from excel_formulas import WorkbookEvaluator
from agents.registry import FINANCIAL_DD, RISK, STRUCTURING
from memo import InvestmentCommitteeMemoGenerator
from models import load_deal
from orchestrator import DealTeamOrchestrator
from test_memo import ScriptedLLM, approved, memo_json, revise, v

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DEAL_PATH = ROOT / "data" / "sample_deal.json"
BLUE, BLACK, GREEN = "000000FF", "00000000", "00008000"
approx = pytest.approx


@pytest.fixture(autouse=True)
def no_api_key(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)


@pytest.fixture(scope="module")
def deal():
    return load_deal(SAMPLE_DEAL_PATH)


def approved_state(deal, ic=None, conclusion="Proceed to further diligence"):
    llm = ScriptedLLM(ic or [approved(gaps=["Customer contracts are not supplied.",
                                            "Financing terms are not supplied."])],
                      memo_response=memo_json(conclusion))
    result = DealTeamOrchestrator(llm).run(deal)
    memo = InvestmentCommitteeMemoGenerator(llm).generate(result, deal)
    return result, memo


@pytest.fixture(scope="module")
def state(deal):
    return approved_state(deal)


@pytest.fixture(scope="module")
def saved(deal, state, tmp_path_factory):
    result, memo = state
    path = export_ma_workbook(deal, result, memo, tmp_path_factory.mktemp("x") / "model.xlsx")
    return path, load_workbook(path)


@pytest.fixture(scope="module")
def values(saved):
    return WorkbookEvaluator(saved[1])


def find(ws, label, column=1):
    for row in ws.iter_rows():
        cell = row[column - 1]
        if cell.value == label:
            return cell.row
    raise AssertionError(f"{label!r} not found in {ws.title}")


def cell_right_of(ws, label):
    return ws.cell(find(ws, label), 2)


def all_text(ws):
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


# --- Eligibility ----------------------------------------------------------------------------

def test_review_limit_reached_cannot_be_exported(deal, tmp_path):
    unresolved = DealTeamOrchestrator(ScriptedLLM([revise(RISK), revise(RISK)])).run(deal)
    with pytest.raises(ExcelExportError, match="unresolved"):
        export_ma_workbook(deal, unresolved, None, tmp_path / "x.xlsx")
    assert not (tmp_path / "x.xlsx").exists()


@pytest.mark.parametrize("change, message", [
    (lambda r: replace(r, status=None), "no final status"),
    (lambda r: replace(r, outputs={k: x for k, x in r.outputs.items() if k != RISK}),
     "missing analytical outputs"),
    (lambda r: replace(r, deterministic_contexts={
        k: x for k, x in r.deterministic_contexts.items() if k != "synergy"}),
     "missing deterministic contexts"),
    (lambda r: replace(r, review_history=[]), "no IC review"),
])
def test_invalid_states_rejected(deal, state, change, message):
    with pytest.raises(ExcelExportError, match=message):
        build_workbook(deal, change(state[0]), None)


def test_result_must_match_the_deal(deal, state):
    other = replace(deal, transaction_assumptions=replace(
        deal.transaction_assumptions, offer_enterprise_value=210.0))
    with pytest.raises(ExcelExportError, match="does not match the deal"):
        build_workbook(other, state[0], None)


def test_memo_must_correspond_to_result(deal, state):
    revised_result, revised_memo = approved_state(deal, ic=[revise(FINANCIAL_DD), approved()])
    with pytest.raises(ExcelExportError, match="does not correspond"):
        build_workbook(deal, state[0], revised_memo)
    with pytest.raises(ExcelExportError, match="approved review"):
        build_workbook(deal, state[0], replace(state[1], orchestration_status="review_limit_reached"))
    with pytest.raises(ExcelExportError, match="different deal"):
        build_workbook(deal, state[0], replace(state[1], deal_name="Other"))


def test_export_does_not_mutate_result(deal, state, tmp_path):
    result, memo = state
    before = (dict(result.outputs), list(result.agent_runs), list(result.review_history),
              result.status)
    export_ma_workbook(deal, result, memo, tmp_path / "m.xlsx")
    assert before == (dict(result.outputs), list(result.agent_runs),
                      list(result.review_history), result.status)


# --- Structure and round-trip -----------------------------------------------------------------

def test_workbook_opens_with_exact_sheet_order(saved):
    assert saved[1].sheetnames == list(SHEET_ORDER)


def test_gridlines_hidden_everywhere(saved):
    assert all(ws.sheet_view.showGridLines is False for ws in saved[1].worksheets)


def test_footer_note_on_every_sheet(saved):
    for ws in saved[1].worksheets:
        assert ws["A2"].value == excel_export.FOOTER_NOTE
        assert "financial advice" not in " ".join(all_text(ws)).lower()


def test_important_headers_present(saved):
    wb = saved[1]
    expected = {
        "Assumptions": ["DEAL OVERVIEW", "ACQUIRER FINANCIALS", "TARGET FINANCIALS",
                        "MARKET VALUATION INPUTS", "DCF ASSUMPTIONS",
                        "SYNERGY ASSUMPTIONS (UNVALIDATED)", "TRANSACTION ASSUMPTIONS"],
        "Valuation": ["VALUATION SUMMARY", "COMPARABLE COMPANIES", "PRECEDENT TRANSACTIONS",
                      "DISCOUNTED CASH FLOW"],
        "Synergies": ["SYNERGY VALUATION", "DISCOUNTED SYNERGY SCHEDULE", "MODELLING CONVENTIONS"],
        "Transaction Structure": ["PURCHASE PRICE", "CONSIDERATION", "FINANCING MODEL CONVENTION",
                                  "SHARE ISSUANCE / OWNERSHIP"],
        "Accretion Dilution": ["CASE A — EXCLUDING SYNERGIES",
                               "CASE B — INCLUDING FULL RUN-RATE SYNERGIES"],
        "Due Diligence": ["FINANCIAL DUE DILIGENCE", "COMMERCIAL DUE DILIGENCE",
                          "LEGAL & OPERATIONAL DUE DILIGENCE"],
        "Risk Summary": ["RISK ANALYST — FINAL REVIEWED OUTPUT", "IC PARTNER — FINAL QUALITY REVIEW"],
        "IC Summary": ["TRANSACTION", "IC CONCLUSION", "KEY UNRESOLVED ITEMS (FINAL IC REVIEW)"],
    }
    for sheet, headers in expected.items():
        text = all_text(wb[sheet])
        for header in headers:
            assert header in text, (sheet, header)


def test_comments_survive_round_trip(saved):
    ws = saved[1]["Assumptions"]
    revenue = ws.cell(find(ws, "Revenue (USD millions)"), 2)
    assert revenue.comment.text == "Source: data/sample_deal.json — acquirer.revenue"
    wacc = cell_right_of(ws, "WACC")
    assert wacc.comment.text == "Source: data/sample_deal.json — dcf_assumptions.wacc"
    dd = saved[1]["Due Diligence"]
    assert any(c.comment and c.comment.text == f"Source: final reviewed output — {FINANCIAL_DD}"
               for row in dd.iter_rows() for c in row)


def test_excel_recalculates_on_open(saved):
    calc = saved[1].calculation  # reloaded from the saved file
    assert calc.calcMode == "auto"
    assert calc.fullCalcOnLoad is True
    assert calc.forceFullCalc is True
    with zipfile.ZipFile(saved[0]) as z:
        xml = z.read("xl/workbook.xml").decode()
    assert 'calcMode="auto"' in xml and 'fullCalcOnLoad="1"' in xml and 'forceFullCalc="1"' in xml


def test_exactly_one_chart_on_valuation(saved):
    with zipfile.ZipFile(saved[0]) as z:
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        assert charts == ["xl/charts/chart1.xml"]
        drawings = [n for n in z.namelist() if n.startswith("xl/drawings/drawing")]
        assert len(drawings) == 1
        sheet_xml = z.read("xl/worksheets/sheet3.xml").decode()  # third sheet: Valuation
        assert "<drawing" in sheet_xml
        chart = z.read("xl/charts/chart1.xml").decode()
        assert "barChart" in chart and "bar3DChart" not in chart


# --- Formulas ----------------------------------------------------------------------------------

def test_no_external_links_or_broken_references(saved):
    wb = saved[1]
    assert wb._external_links == []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    assert "[" not in c.value and "#REF!" not in c.value, (ws.title, c.coordinate)
    WorkbookEvaluator(wb).evaluate_all()  # raises on any unknown sheet or empty reference


def test_key_outputs_are_formulas_not_hardcoded(saved):
    wb = saved[1]
    for sheet, label in [
        ("Valuation", "Implied Enterprise value (USD millions)"),
        ("Valuation", "DCF Enterprise value (USD millions)"),
        ("Synergies", "Net synergy value (USD millions)"),
        ("Transaction Structure", "Offer equity value (USD millions)"),
        ("Transaction Structure", "Cash consideration (USD millions)"),
        ("Transaction Structure", "New shares issued (millions)"),
        ("Accretion Dilution", "Pro forma EPS (USD per share)"),
    ]:
        cell = cell_right_of(wb[sheet], label)
        assert cell.data_type == "f", (sheet, label, cell.value)


def test_formula_font_colours(saved):
    wb = saved[1]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    expected = GREEN if "!" in c.value else BLACK
                    assert c.font.color.rgb == expected, (ws.title, c.coordinate, c.value)
    assert cell_right_of(wb["Valuation"], "Low").font.color.rgb == BLACK
    assert cell_right_of(wb["Transaction Structure"], "Cash %").value.startswith("='Assumptions'!")


def test_assumption_inputs_blue_and_unverified_yellow(saved):
    ws = saved[1]["Assumptions"]
    inputs = [ws.cell(r, 2) for r in range(4, ws.max_row + 1)
              if ws.cell(r, 2).value is not None and ws.cell(r, 1).value is not None]
    assert len(inputs) >= 30
    assert all(c.font.color.rgb == BLUE and c.data_type != "f" for c in inputs)
    for label in ["Annual cost synergy (USD millions)", "Annual revenue synergy EBITDA (USD millions)",
                  "Realization years", "Integration cost (USD millions)", "New debt interest rate",
                  "WACC"]:
        assert cell_right_of(ws, label).fill.fgColor.rgb == "00FFF2CC", label
    assert cell_right_of(ws, "Offer enterprise value (USD millions)").fill.fgColor.rgb != "00FFF2CC"


def test_number_formats_follow_finance_convention(saved):
    wb = saved[1]
    money = cell_right_of(wb["Transaction Structure"], "Offer equity value (USD millions)")
    assert money.number_format == "#,##0.0;[Red](#,##0.0);-"
    assert cell_right_of(wb["Assumptions"], "WACC").number_format == "0.0%;[Red](0.0%);-"
    assert cell_right_of(wb["Valuation"], "Comparable EV / EBITDA (x)").number_format == "0.0x"
    assert cell_right_of(wb["Accretion Dilution"], "Pro forma EPS (USD per share)").number_format \
        == "0.0000;[Red](0.0000);-"
    assert cell_right_of(wb["Transaction Structure"], "Seller ownership").number_format == "0.0%"
    assert cell_right_of(wb["Transaction Structure"], "New shares issued (millions)").number_format \
        == "#,##0.00"


# --- Deterministic consistency (evaluated from the saved file) ---------------------------------

def test_sample_values_from_saved_workbook(saved, values):
    wb = saved[1]

    def val(sheet, label, occurrence=0):
        ws = wb[sheet]
        rows = [c.row for c in ws["A"] if c.value == label]
        return values.value(sheet, f"B{rows[occurrence]}")

    assert val("Transaction Structure", "Target net debt (USD millions)") == approx(20.0)
    assert val("Transaction Structure", "Offer equity value (USD millions)") == approx(185.0)
    assert val("Valuation", "Comparable Companies") == approx(189.0)
    assert val("Valuation", "Precedent Transactions") == approx(216.0)
    assert val("Valuation", "Discounted Cash Flow") == approx(172.024, abs=1e-3)
    assert val("Transaction Structure", "Cash consideration (USD millions)") == approx(111.0)
    assert val("Transaction Structure", "Stock consideration (USD millions)") == approx(74.0)
    assert val("Transaction Structure", "New shares issued (millions)") == approx(2.96)
    assert val("Transaction Structure", "Pro forma shares (millions)") == approx(102.96)
    assert val("Synergies", "PV of synergies (USD millions)") == approx(25.5878, abs=1e-4)
    assert val("Synergies", "Net synergy value (USD millions)") == approx(13.5878, abs=1e-4)
    assert val("Accretion Dilution", "Result", 0) == "Accretive"
    assert val("Accretion Dilution", "Result", 1) == "Accretive"


def test_every_formula_matches_python_full_precision(deal, saved, values):
    """Independent re-check on the saved file, against the deterministic contexts."""
    from deal_analysis import build_transaction_structure_context, build_valuation_context
    val = build_valuation_context(deal, ndigits=None)
    tx = build_transaction_structure_context(deal, ndigits=None)
    ws = saved[1]["Valuation"]
    first_year_row = find(ws, "Year") + 1
    for i, expected in enumerate(val["dcf"]["discounted_fcfs"]):
        assert values.value("Valuation", f"D{first_year_row + i}") == approx(expected, rel=1e-12)
    assert values.value("Valuation", f"B{find(ws, 'Midpoint')}") == approx(
        val["valuation_summary"]["midpoint"], rel=1e-12)
    acc = saved[1]["Accretion Dilution"]
    pf_rows = [c.row for c in acc["A"] if c.value == "Pro forma EPS (USD per share)"]
    assert values.value("Accretion Dilution", f"B{pf_rows[1]}") == approx(
        tx["calculation"]["accretion_dilution_including_run_rate_synergies"]["pro_forma_eps"],
        rel=1e-12)


def test_validation_rejects_workbook_that_disagrees_with_python(deal, state, monkeypatch):
    real = excel_export.build_transaction_structure_context

    def skewed(d, ndigits=4):
        ctx = real(d, ndigits=ndigits)
        if ndigits is None:  # only the full-precision copy used for formula checks
            ctx["calculation"]["cash_consideration"] = 999.0
        return ctx

    monkeypatch.setattr(excel_export, "build_transaction_structure_context", skewed)
    with pytest.raises(ExcelExportError, match="disagrees with deterministic Python"):
        build_workbook(deal, state[0], state[1])


# --- Content rules ------------------------------------------------------------------------------

def test_valuation_has_three_methods_and_labels(saved):
    text = all_text(saved[1]["Valuation"])
    for method in ["Comparable Companies", "Precedent Transactions", "Discounted Cash Flow"]:
        assert method in text
    assert "Offer vs range: within" in text
    assert "premium" not in " ".join(text).lower().replace("no premium or discount", "")


def test_synergy_schedule_and_conventions(saved):
    ws = saved[1]["Synergies"]
    first = find(ws, "Year") + 1
    assert [ws.cell(first + i, 1).value for i in range(5)] == [1, 2, 3, 4, 5]
    text = " ".join(all_text(ws))
    for convention in ["Full run-rate synergies from year one", "no terminal synergy value",
                       "Integration cost is incurred at close", "unvalidated"]:
        assert convention in text


def test_financing_convention_warning(saved):
    text = all_text(saved[1]["Transaction Structure"])
    assert FINANCING_WARNING in text
    assert any("funded entirely with new debt" in t for t in text)
    joined = " ".join(text).lower()
    for absent in ["covenant headroom", "debt capacity"]:
        assert absent in joined  # only in the "not modelled" note
    assert "Not modelled: leverage, debt capacity, covenant headroom, or credit-rating impact." in text


def test_accretion_both_cases_and_caution(saved):
    text = all_text(saved[1]["Accretion Dilution"])
    assert "CASE A — EXCLUDING SYNERGIES" in text
    assert "CASE B — INCLUDING FULL RUN-RATE SYNERGIES" in text
    assert SYNERGY_CASE_CAUTION in text
    for workbook_text in (all_text(ws) for ws in saved[1].worksheets):
        for banned in ["Base Case", "Expected Case", "Most Likely Case"]:
            assert not any(banned in t for t in workbook_text)


def test_no_scores_or_ratings(saved):
    for ws in saved[1].worksheets:
        text = " ".join(all_text(ws))
        for banned in ["High / Medium / Low", "Probability", "Score", "Traffic light"]:
            assert banned not in text


# --- Current vs stale outputs --------------------------------------------------------------------

def test_current_versions_used_and_stale_excluded(deal, tmp_path):
    result, memo = approved_state(deal, ic=[revise(FINANCIAL_DD), approved()])
    wb = load_workbook(export_ma_workbook(deal, result, memo, tmp_path / "r.xlsx"))
    dd = all_text(wb["Due Diligence"])
    assert v(FINANCIAL_DD, 1) in dd and v(FINANCIAL_DD, 0) not in dd
    assert any(t.startswith("Current version: 2 — run kind: direct_revision") for t in dd)
    risk = all_text(wb["Risk Summary"])
    assert v(RISK, 1) in risk and v(RISK, 0) not in risk
    assert "Review round: 2 of 2 used" in risk
    everything = [t for ws in wb.worksheets for t in all_text(ws)]
    assert v(STRUCTURING, 0) not in everything


def test_agent_text_is_never_turned_into_a_formula(deal, tmp_path):
    result, memo = approved_state(deal)
    tampered = replace(result, outputs={**result.outputs, FINANCIAL_DD: "=SUM(1,1)\nplain line"})
    runs = [r if r.agent != FINANCIAL_DD else replace(r, output=tampered.outputs[FINANCIAL_DD])
            for r in result.agent_runs]
    tampered = replace(tampered, agent_runs=runs)
    wb = load_workbook(export_ma_workbook(deal, tampered, None, tmp_path / "t.xlsx"))
    cell = next(c for row in wb["Due Diligence"].iter_rows() for c in row if c.value == "=SUM(1,1)")
    assert cell.data_type == "s"


# --- IC Summary / memo -----------------------------------------------------------------------------

def test_ic_summary_uses_final_review_and_memo(saved, state):
    result, memo = state
    text = all_text(saved[1]["IC Summary"])
    assert memo.conclusion in text
    assert "• Customer contracts are not supplied." in text
    assert "• Financing terms are not supplied." in text
    ws = saved[1]["IC Summary"]
    assert cell_right_of(ws, "IC review status").value == "approved"
    assert cell_right_of(ws, "Review rounds used").value == 1


def test_memo_sheet_contains_full_memo(saved, state):
    memo = state[1]
    text = all_text(saved[1]["IC Memo"])
    assert text[2] == "IC CONCLUSION" and text[3] == memo.conclusion
    for line in memo.memo_text.splitlines():
        if line.strip():
            assert line.strip().lstrip("#").strip() in text


def test_without_memo(deal, state, tmp_path):
    wb = load_workbook(export_ma_workbook(deal, state[0], None, tmp_path / "n.xlsx"))
    assert NO_MEMO_NOTE in all_text(wb["IC Memo"])
    assert "No IC memo supplied" in all_text(wb["IC Summary"])
    assert wb.sheetnames == list(SHEET_ORDER)


# --- Files / dependencies ------------------------------------------------------------------------

def test_overwrite_protection(deal, state, tmp_path):
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"unrelated")
    with pytest.raises(ExcelExportError, match="already exists"):
        export_ma_workbook(deal, state[0], state[1], target)
    assert target.read_bytes() == b"unrelated"
    export_ma_workbook(deal, state[0], state[1], target, overwrite=True)
    assert load_workbook(target).sheetnames == list(SHEET_ORDER)
    with pytest.raises(ExcelExportError, match=".xlsx"):
        export_ma_workbook(deal, state[0], state[1], tmp_path / "model.csv")


def test_exporter_has_no_llm_dependency():
    tree = ast.parse((ROOT / "excel_export.py").read_text())
    imported = {alias.name for node in ast.walk(tree) if isinstance(node, ast.Import)
                for alias in node.names}
    imported |= {node.module for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)}
    assert not any(m and (m.startswith("google") or m == "llm_client" or m == "config")
                   for m in imported)
    source = (ROOT / "excel_export.py").read_text()
    assert ".generate(" not in source and "GeminiClient" not in source


def test_sample_script_works_offline(tmp_path, monkeypatch):
    import export_sample_workbook
    monkeypatch.setattr("sys.argv", ["export_sample_workbook.py", str(tmp_path / "s.xlsx")])
    assert export_sample_workbook.main() == 0
    wb = load_workbook(tmp_path / "s.xlsx")
    assert wb.sheetnames == list(SHEET_ORDER)
    assert any("Offline Demo" in t for t in all_text(wb["Due Diligence"]))
